import os
import pytest
import openpyxl
from app.finance.dcf.xlsx_writer import write_dcf_xlsx
from app.finance.dcf.assumptions import defaults_from_ssot
from app.finance.dcf.run import run_dcf

def test_xlsx_writer(tmp_path):
    a = defaults_from_ssot([], [], "2026-10-23")
    res = run_dcf(a)
    
    file_path = tmp_path / "ONGC_DCF_test_run.xlsx"
    write_dcf_xlsx(res, str(file_path))
    
    assert file_path.exists()
    
    wb = openpyxl.load_workbook(str(file_path))
    assert "Cover" in wb.sheetnames
    assert "Assumptions" in wb.sheetnames
    assert "WACC" in wb.sheetnames
    assert "Historical" in wb.sheetnames
    assert "Forecast" in wb.sheetnames
    assert "Valuation" in wb.sheetnames
    assert "Sensitivity" in wb.sheetnames
    assert "Tornado" in wb.sheetnames
    assert "Sources" in wb.sheetnames
    
    ws_v = wb["Valuation"]
    val_formula = ws_v["B15"].value
    assert isinstance(val_formula, str) and val_formula.startswith("=")
    
    # Check that named ranges were created for assumptions
    assert "wacc_pct" in wb.defined_names
    
    # Check cover bold disclaimer
    assert wb["Cover"]["A3"].font.bold

    # Add pure-Python formula parity check
    from tests.finance.evaluator import evaluate_dcf_excel
    
    # 1. Valuation B15 parity
    calc_val = evaluate_dcf_excel(wb, "Valuation!B15")
    py_val = res["per_share"].get("intrinsic_inr")
    assert abs(calc_val - py_val) < 0.01, f"Parity mismatch: {calc_val} vs {py_val}"
    
    # 2. Sensitivity parity
    s_c5 = evaluate_dcf_excel(wb, "Sensitivity!C5")
    s_f7 = evaluate_dcf_excel(wb, "Sensitivity!F7")
    s_j9 = evaluate_dcf_excel(wb, "Sensitivity!J9")
    
    grid = res["sensitivity"]["values_inr"]
    assert abs(s_c5 - grid[2][1]) < 0.01, f"Sensitivity C5 mismatch"
    assert abs(s_f7 - grid[4][4]) < 0.01, f"Sensitivity F7 mismatch"
    assert abs(s_j9 - grid[6][8]) < 0.01, f"Sensitivity J9 mismatch"
    
    # 3. Proving wacc_pct changes B15
    calc_val_new = evaluate_dcf_excel(wb, "Valuation!B15", wacc_override=12.0)
    assert abs(calc_val_new - calc_val) > 1.0, "wacc_pct override failed"
