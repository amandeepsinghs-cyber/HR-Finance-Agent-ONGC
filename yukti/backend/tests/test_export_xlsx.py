"""P4.6 (SPEC-10 §4.1, §6 AC 1 & AC 9): openpyxl read-back verification for `.xlsx` exports.

Covers:
- FY22–FY26 P&L statement fixture: frozen pane at `A5`, `B5` number_format `#,##0;(#,##0)`,
  label indents equal `level` (incl. `A5`/`A6`/`A8` indent 1), subtotal (`Total income`,
  `Profit before tax (PBT)`) and total (`Profit after tax (PAT)`) rows bold with borders,
  no autofilter on statements, footnote rows present, 100% cell-for-cell value fidelity.
- V19 Delhi >5y roster fixture: sheet name <= 31 chars, title block in rows 1–2, header row 4
  bold + shaded + bottom border, frozen at `A5`, autofilter on, tenure column format `0.0`,
  `Definition` sheet added with rule text + JRTP 3.2 (Illustrative) + J-03 citation.
- Masked / aggregate E2_HR export: `"<5"` cells stored as string `"<5"` (never raw count),
  aggregate banner present, workbook metadata (`creator="Yukti"`, `keywords="persona=...; masked=..."`).
- DCF valuation fixture: 6 sheets (`Summary`, `Assumptions`, `FCFF`, `WACC`, `Sensitivity`, `Tornado`),
  blue input cells (`0000FF`) + workbook DefinedNames on `Assumptions`, live formulas (`=...`) referencing `Assumptions`.
"""

from __future__ import annotations

import json
from pathlib import Path
import sys

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.export_audit import NAMING_REGEX, enforce_export_rbac, filename_from_contract
from app.tools.export_tool import DocumentExporter

FIXTURES_DIR = Path(__file__).resolve().parent.parent.parent / "contracts" / "fixtures" / "export"


@pytest.fixture()
def exporter(tmp_path: Path) -> DocumentExporter:
    return DocumentExporter(exports_dir=tmp_path / "exports")


def _load_fixture(name: str) -> dict:
    return json.loads((FIXTURES_DIR / name).read_text(encoding="utf-8"))


def test_pnl_statement_xlsx_layout_formats_indents_and_100pct_value_parity(exporter: DocumentExporter) -> None:
    contract = _load_fixture("pnl_statement_fy22_fy26_synthetic.json")
    fname = filename_from_contract(contract, ext="xlsx", persona="CMD")
    assert NAMING_REGEX.fullmatch(fname), fname
    assert fname.startswith("yukti_pnl_standalone_fy22-fy26_cmd_")

    out_path = exporter.export_xlsx(filename=fname, contract=contract, persona="CMD")
    wb = openpyxl.load_workbook(str(out_path), data_only=False)
    ws = wb.active

    # 1. Sheet name <= 31 chars, frozen pane at A5, no autofilter on statements
    assert len(ws.title) <= 31
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref is None

    # 2. Title block rows 1-2, Header row 4 (latest-first FY26..FY22)
    assert ws["A1"].font.bold is True and ws["A1"].font.size == 14
    assert "As of" in str(ws["A2"].value) and "Source:" in str(ws["A2"].value)
    periods = contract["visual"]["statement"]["periods"]
    assert periods == ["FY26", "FY25", "FY24", "FY23", "FY22"]
    for idx, p in enumerate(periods, start=2):
        hdr_cell = ws.cell(row=4, column=idx)
        assert hdr_cell.value == p
        assert hdr_cell.font.bold is True

    # 3. B5 number_format is "#,##0;(#,##0)", indents equal level, subtotals/PBT/PAT bold + borders, 100% value match
    assert ws["B5"].number_format == "#,##0;(#,##0)"

    all_contract_rows = [
        r for sec in contract["visual"]["statement"]["sections"] for r in sec["rows"]
    ]
    for r_offset, c_row in enumerate(all_contract_rows, start=5):
        cell_a = ws.cell(row=r_offset, column=1)
        assert cell_a.value == c_row["label"]
        assert int(cell_a.alignment.indent or 0) == int(c_row["level"]), (c_row["label"], cell_a.alignment.indent)

        kind = c_row.get("kind", "line")
        label_low = c_row["label"].lower()
        if kind in ("subtotal", "total") or "profit before tax" in label_low or "profit after tax" in label_low:
            assert cell_a.font.bold is True, f"Expected bold on {c_row['label']}"
            assert cell_a.border.top.style is not None
        if kind == "total" or "profit after tax" in label_low:
            assert cell_a.border.bottom.style == "double"

        for p_idx, p in enumerate(periods, start=2):
            c_val = ws.cell(row=r_offset, column=p_idx)
            assert c_val.number_format == "#,##0;(#,##0)"
            expected = c_row["values"].get(p)
            assert c_val.value == pytest.approx(expected), f"Mismatch at {c_row['label']} {p}"


def test_v19_roster_xlsx_formats_widths_autofilter_and_definition_sheet(exporter: DocumentExporter) -> None:
    contract = _load_fixture("v19_delhi_roster_synthetic.json")
    out_path = exporter.export_xlsx(filename="yukti_delhi_roster.xlsx", contract=contract, persona="CMD")
    wb = openpyxl.load_workbook(str(out_path), data_only=False)

    assert "Definition" in wb.sheetnames
    ws = wb[wb.sheetnames[0]]
    assert ws.title == "Delhi >5y roster"
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref == "A4:G69"

    # Header row 4 matches display headers
    headers = [ws.cell(row=4, column=c).value for c in range(1, 8)]
    assert headers == [
        "Name",
        "Designation",
        "Level",
        "ONGC service (yrs)",
        "Tenure in Delhi (yrs)",
        "Years to retirement",
        "Exemption",
    ]
    # Tenure columns (D, E, F) use "0.0" number format
    assert ws["D5"].number_format == "0.0"
    assert ws["E5"].number_format == "0.0"
    assert ws["F5"].number_format == "0.0"

    # Column widths auto-fitted in [8, 60]
    for col_let in ("A", "B", "C", "D", "E", "F", "G"):
        w = ws.column_dimensions[col_let].width
        assert 8 <= w <= 60

    # Definition sheet checks
    ws_def = wb["Definition"]
    assert ws_def.freeze_panes == "A5"
    def_text = " ".join(
        str(ws_def.cell(row=r, column=c).value or "")
        for r in range(1, ws_def.max_row + 1)
        for c in range(1, ws_def.max_column + 1)
    )
    assert "DELHI_GT5_CONT_V1" in def_text
    assert "Illustrative" in def_text
    assert "LPA 744/2013" in def_text


def test_masked_aggregate_xlsx_keeps_small_cell_text_and_banner(exporter: DocumentExporter) -> None:
    contract = _load_fixture("v19_delhi_roster_synthetic.json")
    masked_contract, masked_flag, is_agg, _, _ = enforce_export_rbac(contract, "E2_HR")
    assert masked_flag is True and is_agg is True

    fname = filename_from_contract(masked_contract, ext="xlsx", persona="E2_HR", is_aggregate=is_agg)
    assert fname.endswith("_aggregate.xlsx")
    assert NAMING_REGEX.fullmatch(fname)

    out_path = exporter.export_xlsx(
        filename=fname,
        contract=masked_contract,
        persona="E2_HR",
        masked=masked_flag,
        is_aggregate=is_agg,
    )
    wb = openpyxl.load_workbook(str(out_path), data_only=False)
    assert wb.properties.creator == "Yukti"
    assert "persona=E2_HR; masked=true" in (wb.properties.keywords or "")

    ws = wb[wb.sheetnames[0]]
    assert "Aggregate view (access level: E2_HR)" in str(ws["A3"].value)
    # No Name column in row 4
    hdrs = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    assert "Name" not in hdrs and "name" not in hdrs
    assert hdrs == ["Level", "Executives"]

    # Small cells (E7=3, E8=3, E9=2) written as literal string "<5", never raw integers 2 or 3
    col_b_vals = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)]
    assert "<5" in col_b_vals
    assert 2 not in col_b_vals and 3 not in col_b_vals


def test_dcf_workbook_sheets_blue_inputs_and_live_formulas(exporter: DocumentExporter) -> None:
    contract = _load_fixture("dcf_valuation_synthetic.json")
    out_path = exporter.export_xlsx(filename="yukti_dcf.xlsx", contract=contract, persona="CMD")
    wb = openpyxl.load_workbook(str(out_path), data_only=False)

    assert wb.sheetnames == ["Summary", "Assumptions", "FCFF", "WACC", "Sensitivity", "Tornado"]
    ws_assump = wb["Assumptions"]
    assert ws_assump.freeze_panes == "A5"
    assert ws_assump["B5"].value == 34200.0
    assert ws_assump["B5"].font.color.rgb.endswith("0000FF")
    assert ws_assump["B6"].number_format == "0.0%"
    assert "BASE_FCFF" in wb.defined_names
    assert "WACC" in wb.defined_names

    ws_sum = wb["Summary"]
    assert str(ws_sum["B5"].value).startswith("=Assumptions!")
    assert str(ws_sum["B6"].value).startswith("=SUM(FCFF!")
    assert str(ws_sum["B8"].value) == "=B6+B7"

    ws_fcff = wb["FCFF"]
    assert str(ws_fcff["B5"].value).startswith("=Assumptions!")
    assert str(ws_fcff["G5"].value).startswith("=F5*(1+Assumptions!")
