"""P4.6 (SPEC-10 §4.2, §4.5, §4.6, §6 AC 2 & AC 5): CSV, Word (`.docx`), and Email (`.eml`) export tests.

Verifies:
- CSV begins with UTF-8 BOM (`EF BB BF`), preserves `₹` and Devanagari intact, quotes commas/newlines
  per RFC 4180, adds `(%)` suffix to percent headers, and flattens statements with `level` and `kind` columns.
- Multi-table compare contract exports as a `.zip` of UTF-8 BOM CSVs.
- `.docx` round-trips cleanly via `python-docx` with table + embedded `chart_png` image (`inline_shapes`),
  statement indentation (`left_indent`), and email draft letter layout.
- `.eml` parses via Python `email.message_from_bytes` with `X-Unsent: 1`, UTF-8 Hindi subject,
  `text/plain` + `text/html` body parts, and 1 `.xlsx` attachment (`application/vnd.openxmlformats-officedocument.spreadsheetml.sheet`).
"""

from __future__ import annotations

import base64
import csv
import email
from email import policy
import io
import json
from pathlib import Path
import sys
import zipfile

import docx
import openpyxl
from PIL import Image
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.tools.export_tool import DocumentExporter

FIXTURES_DIR = Path(__file__).resolve().parent.parent.parent / "contracts" / "fixtures" / "export"


@pytest.fixture()
def exporter(tmp_path: Path) -> DocumentExporter:
    return DocumentExporter(exports_dir=tmp_path / "exports")


def _load_fixture(name: str) -> dict:
    return json.loads((FIXTURES_DIR / name).read_text(encoding="utf-8"))


def _make_sample_png_data_url() -> str:
    buf = io.BytesIO()
    Image.new("RGB", (320, 160), color=(20, 80, 140)).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


def test_csv_utf8_bom_devanagari_rupee_and_quoting(exporter: DocumentExporter) -> None:
    hi = _load_fixture("hindi_revenue_pat_synthetic.json")
    csv_path = exporter.export_csv(filename="yukti_hindi.csv", contract=hi)

    raw_bytes = csv_path.read_bytes()
    # AC 2: First 3 bytes must be UTF-8 BOM EF BB BF
    assert raw_bytes[:3] == b"\xef\xbb\xbf"

    decoded = raw_bytes.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(decoded)))
    assert rows[0] == ["वित्त वर्ष (FY)", "राजस्व (₹ करोड़)", "शुद्ध लाभ PAT (₹ करोड़)", "PAT मार्जिन (%)"]
    assert rows[1][0] == "FY26"
    assert float(rows[1][1]) == pytest.approx(132508.14)
    assert float(rows[1][2]) == pytest.approx(38348.07)

    # RFC 4180 quoting with commas and negatives
    custom_path = exporter.export_csv(
        columns=["label", "delta_cr", "margin_pct"],
        rows=[{"label": "Oil, Gas & Condensate (₹)", "delta_cr": -425.5, "margin_pct": 18.4}],
        filename="yukti_quoting.csv",
    )
    q_bytes = custom_path.read_bytes()
    assert q_bytes[:3] == b"\xef\xbb\xbf"
    q_rows = list(csv.reader(io.StringIO(q_bytes.decode("utf-8-sig"))))
    assert q_rows[0] == ["label", "delta_cr", "margin_pct (%)"]
    assert q_rows[1] == ["Oil, Gas & Condensate (₹)", "-425.5", "18.4"]


def test_csv_flattens_statement_with_level_and_kind_and_zips_multi_table(exporter: DocumentExporter) -> None:
    pnl = _load_fixture("pnl_statement_fy22_fy26_synthetic.json")
    csv_path = exporter.export_csv(filename="yukti_pnl.csv", contract=pnl)
    raw = csv_path.read_bytes()
    assert raw[:3] == b"\xef\xbb\xbf"

    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0][:4] == ["line_code", "label", "level", "kind"]
    assert "FY26" in rows[0] and "FY22" in rows[0]
    rev_row = next(r for r in rows[1:] if r[0] == "REV_OPS")
    assert rev_row[1] == "Revenue from operations"
    assert rev_row[2] == "1"
    assert rev_row[3] == "line"
    assert float(rev_row[4]) == pytest.approx(132508.14)

    # Multi-table compare -> zip of UTF-8 BOM CSVs
    multi_contract = {
        "visual": {
            "type": "compare",
            "compare": {
                "items": [
                    {"columns": ["wc", "hc"], "rows": [{"wc": "WC-DEL", "hc": 65}]},
                    {"columns": ["wc", "hc"], "rows": [{"wc": "WC-MUM", "hc": 120}]},
                ]
            },
        }
    }
    zip_path = exporter.export_csv(filename="yukti_compare.zip", contract=multi_contract)
    assert zip_path.suffix == ".zip"
    with zipfile.ZipFile(zip_path, "r") as zf:
        assert sorted(zf.namelist()) == ["table_1.csv", "table_2.csv"]
        assert zf.read("table_1.csv")[:3] == b"\xef\xbb\xbf"


def test_docx_roundtrip_table_image_statement_and_email_letter(exporter: DocumentExporter) -> None:
    pnl = _load_fixture("pnl_statement_fy22_fy26_synthetic.json")
    chart_png = _make_sample_png_data_url()
    docx_path = exporter.export_word(
        filename="yukti_pnl.docx",
        contract=pnl,
        chart_png=chart_png,
        persona="CMD",
    )

    doc = docx.Document(str(docx_path))
    assert len(doc.tables) == 1
    assert len(doc.inline_shapes) == 1
    tbl = doc.tables[0]
    assert tbl.cell(0, 0).text == "Line Item (₹ Crore)"
    assert tbl.cell(0, 1).text == "FY26"
    # Revenue from operations has level 1 -> 12pt left_indent
    assert tbl.cell(1, 0).text == "Revenue from operations"
    assert tbl.cell(1, 0).paragraphs[0].paragraph_format.left_indent == docx.shared.Pt(12)

    # Email draft letter mode
    hi = _load_fixture("hindi_revenue_pat_synthetic.json")
    letter_path = exporter.export_word(
        filename="yukti_letter.docx",
        contract=hi,
        email_draft=hi["email_draft"],
        persona="CMD",
    )
    letter_doc = docx.Document(str(letter_path))
    full_text = "\n".join(p.text for p in letter_doc.paragraphs)
    assert "Director (HR)" in full_text
    assert "समीक्षा अनुरोध" in full_text
    assert len(letter_doc.tables) == 1


def test_eml_parses_with_x_unsent_utf8_headers_and_xlsx_attachment(exporter: DocumentExporter) -> None:
    hi = _load_fixture("hindi_revenue_pat_synthetic.json")
    eml_path = exporter.export_eml(
        email_draft=hi["email_draft"],
        filename="yukti_transfer_draft.eml",
        contract=hi,
        persona="CMD",
    )

    msg = email.message_from_bytes(eml_path.read_bytes(), policy=policy.default)
    assert msg["X-Unsent"] == "1"
    assert "Director (HR)" in msg["To"]
    assert "समीक्षा अनुरोध" in msg["Subject"]

    attachments = list(msg.iter_attachments())
    assert len(attachments) == 1
    att = attachments[0]
    assert att.get_content_type() == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    att_bytes = att.get_payload(decode=True)
    assert len(att_bytes) > 1000

    # Verify the attached xlsx opens cleanly in openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(att_bytes))
    assert len(wb.sheetnames) >= 1
