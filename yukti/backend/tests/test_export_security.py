"""H1.3a + P4.6 (SPEC-10 §4.7, §4.9, §4.10, §6 AC 6, 8, 9):
Export path-traversal, download hardening, 11-persona V19 RBAC matrix, cross-principal 403,
24h expiry, honest Sheets/Docs toasts (no `docs.google.com`), naming regex, and JSONL audit logging.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.export_audit import (
    NAMING_REGEX,
    get_artifact_record,
    read_export_audit_rows,
    register_turn_contract,
)
from app.tools.export_tool import DOCS_TOAST, SHEETS_TOAST, DocumentExporter

FIXTURES_DIR = Path(__file__).resolve().parent.parent.parent / "contracts" / "fixtures" / "export"

ALL_11_PERSONAS = {
    "CMD": {"emp_id": "EMP-000001", "display_name": "Arun Mehra", "designation": "CMD", "function": "EXEC", "grade": "CMD"},
    "DIR_HR": {"emp_id": "EMP-000002", "display_name": "Manish Deshpande", "designation": "Director (HR)", "function": "HR", "grade": "DIR"},
    "DIR_FIN": {"emp_id": "EMP-000003", "display_name": "Vivek Rao", "designation": "Director (Finance)", "function": "FIN", "grade": "DIR"},
    "E1_HR": {"emp_id": "EMP-100151", "display_name": "Dinesh Deshmukh", "designation": "HR Officer (E1)", "function": "HR", "grade": "E1"},
    "E2_HR": {"emp_id": "EMP-102451", "display_name": "Pooja Verma", "designation": "HR Executive (E2)", "function": "HR", "grade": "E2"},
    "E3_HR": {"emp_id": "EMP-100144", "display_name": "Pooja Pillai", "designation": "Senior HR Executive (E3)", "function": "HR", "grade": "E3"},
    "E4_HR": {"emp_id": "EMP-100178", "display_name": "Aarti Saikia", "designation": "Dy. Manager HR (E4)", "function": "HR", "grade": "E4"},
    "E1_FIN": {"emp_id": "EMP-100150", "display_name": "Ajay Deshmukh", "designation": "Accounts Officer (E1)", "function": "FIN", "grade": "E1"},
    "E2_FIN": {"emp_id": "EMP-100152", "display_name": "Praveen Nair", "designation": "Finance Executive (E2)", "function": "FIN", "grade": "E2"},
    "E3_FIN": {"emp_id": "EMP-100146", "display_name": "Asha Mishra", "designation": "Finance Officer (E3)", "function": "FIN", "grade": "E3"},
    "E4_FIN": {"emp_id": "EMP-100145", "display_name": "Santosh Banerjee", "designation": "Dy. Manager Finance (E4)", "function": "FIN", "grade": "E4"},
}


def _auth(persona: str = "CMD", monkeypatch: pytest.MonkeyPatch | None = None) -> dict:
    """P0.4: export routes require a signed session token."""
    from app import auth

    if persona not in auth.PERSONA_DIRECTORY and monkeypatch is not None:
        monkeypatch.setitem(auth.PERSONA_DIRECTORY, persona, ALL_11_PERSONAS[persona])
    return {"Authorization": f"Bearer {auth.issue_token(persona)}"}


@pytest.fixture()
def exporter(tmp_path: Path) -> DocumentExporter:
    return DocumentExporter(exports_dir=tmp_path / "exports")


@pytest.mark.parametrize(
    "name",
    ["../../etc/passwd", "/etc/passwd", "..\\..\\win.ini", "a/../../b.docx", ".bashrc", "", "....//....//x.pdf", "ok name (1).docx"],
)
def test_safe_path_never_escapes(exporter: DocumentExporter, name: str) -> None:
    p = exporter.safe_path(name, "docx")
    assert exporter.exports_dir.resolve() in p.parents
    assert not p.name.startswith(".")


def test_word_export_with_traversal_title_stays_inside(exporter: DocumentExporter) -> None:
    out = exporter.export_word("# hi", filename="../../escape.docx")
    assert out.parent == exporter.exports_dir.resolve()
    assert out.exists()


def test_download_rejects_non_minted_ids() -> None:
    from app.main import app

    c = TestClient(app, headers=_auth())
    for bad in ["..", "%2e%2e", ".gitkeep", "exp_1_abc", "anything", "exp_1792656351753_b136f6", "exp_" + "a" * 32 + "/../x"]:
        r = c.get(f"/api/export/{bad}")
        assert r.status_code == 404, bad


def test_create_then_download_roundtrip() -> None:
    from app.main import app

    c = TestClient(app, headers=_auth())
    r = c.post("/api/export", json={"format": "pdf", "title": "../../Board Note", "content": "Hello"})
    assert r.status_code == 200
    j = r.json()
    assert "/" not in j["filename"] and ".." not in j["filename"]
    assert NAMING_REGEX.fullmatch(j["filename"]), j["filename"]
    d = c.get(f"/api/export/{j['export_id']}")
    assert d.status_code == 200
    assert d.headers["content-type"].startswith("application/pdf")


def test_export_ids_are_128_bit_and_extension_is_forced(tmp_path: Path) -> None:
    from app.main import app

    c = TestClient(app, headers=_auth())
    j = c.post("/api/export", json={"format": "xlsx", "title": "t", "columns": ["a"], "rows": [{"a": 1}]}).json()
    assert len(j["export_id"]) == 4 + 32
    ex = DocumentExporter(exports_dir=tmp_path / "e")
    assert ex.export_word("# x", filename="x.pdf").suffix == ".docx"
    assert ex.export_xlsx(columns=["a"], rows=[{"a": 1}], filename="noext").suffix == ".xlsx"


def test_old_exports_are_swept(monkeypatch) -> None:
    import os, time
    import app.main as m

    import pathlib, tempfile

    tmp = pathlib.Path(tempfile.mkdtemp())
    monkeypatch.setattr(m, "EXPORTS_DIR", tmp)
    old = m.EXPORTS_DIR / "exp_old_file_for_sweep.pdf"
    old.write_bytes(b"x")
    os.utime(old, (time.time() - 7200, time.time() - 7200))
    m._sweep_old_exports()
    assert not old.exists()


def test_errors_do_not_leak_internals(monkeypatch) -> None:
    import app.main as m

    def boom(**_):
        raise ValueError("/secret/path/here")

    monkeypatch.setattr(m.exporter, "export_pdf", boom)
    r = TestClient(m.app, headers=_auth()).post("/api/export", json={"format": "pdf", "title": "t", "content": "x"})
    assert r.status_code == 500 and "/secret" not in r.text


# ------------------------------------------------------------------ P4.6 extensions (SPEC-10 §4.7, §4.9, §4.10)
@pytest.mark.parametrize("persona", list(ALL_11_PERSONAS.keys()))
def test_11_personas_v19_roster_export_rbac_and_audit(persona: str, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from app.main import app

    audit_file = tmp_path / "export_audit.jsonl"
    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(audit_file))

    v19 = json.loads((FIXTURES_DIR / "v19_delhi_roster_synthetic.json").read_text(encoding="utf-8"))
    turn_id = f"turn_v19_{persona.lower()}"
    emp_id = ALL_11_PERSONAS[persona]["emp_id"]
    # Ensure auth directory has the persona with its canonical emp_id
    from app import auth

    monkeypatch.setitem(auth.PERSONA_DIRECTORY, persona, ALL_11_PERSONAS[persona])
    register_turn_contract(turn_id=turn_id, contract=v19, emp_id=emp_id, persona=persona, session_id=f"sess_{persona}")

    client = TestClient(app, headers=_auth(persona, monkeypatch))
    resp = client.post("/api/export", json={"turn_id": turn_id, "format": "xlsx"})
    assert resp.status_code == 200, (persona, resp.text)
    data = resp.json()
    artifact = data["artifact"]

    # Naming regex AC 9
    assert NAMING_REGEX.fullmatch(artifact["filename"]), artifact["filename"]
    stem = artifact["filename"].rsplit(".", 1)[0]
    assert len(stem) <= 80

    # Download file as creator principal
    dl = client.get(f"/api/export/{artifact['export_id']}")
    assert dl.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(dl.content))
    ws = wb[wb.sheetnames[0]]
    hdrs = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]

    if persona in ("CMD", "DIR_HR"):
        assert artifact["masked"] is False
        assert not artifact["filename"].endswith("_aggregate.xlsx")
        assert "Name" in hdrs
        assert artifact["rows"] == 65
    else:
        # DIR_FIN and all 8 juniors (E1_HR..E4_HR, E1_FIN..E4_FIN) get aggregate xlsx with _aggregate suffix
        assert artifact["masked"] is True
        assert artifact["filename"].endswith("_aggregate.xlsx"), artifact["filename"]
        assert "Name" not in hdrs and "name" not in hdrs
        assert f"Aggregate view (access level: {persona})" in str(ws["A3"].value)
        if persona != "DIR_FIN":
            # Junior personas also have small cells (<5) suppressed
            vals = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)]
            assert "<5" in vals

    # Every export writes 1 audit row (100%)
    audit_rows = read_export_audit_rows(audit_file)
    assert len(audit_rows) == 1
    ar = audit_rows[0]
    assert ar["persona"] == persona
    assert ar["emp_id"] == emp_id
    assert ar["turn_id"] == turn_id
    assert ar["format"] == "xlsx"
    assert ar["filename"] == artifact["filename"]
    assert ar["sha256"] == artifact["sha256"]
    assert ar["masked"] is artifact["masked"]


def test_cross_principal_turn_export_and_download_return_403_and_24h_expiry_returns_404(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from app.main import app

    audit_file = tmp_path / "export_audit.jsonl"
    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(audit_file))

    v19 = json.loads((FIXTURES_DIR / "v19_delhi_roster_synthetic.json").read_text(encoding="utf-8"))
    register_turn_contract(
        turn_id="turn_cmd_owned",
        contract=v19,
        emp_id="DEMO-CMD",
        persona="CMD",
        session_id="sess_cmd",
    )

    cmd_client = TestClient(app, headers=_auth("CMD"))
    e2_client = TestClient(app, headers=_auth("E2_HR"))

    # 1. E2_HR trying to export CMD's turn_id -> 403
    r_turn_403 = e2_client.post("/api/export", json={"turn_id": "turn_cmd_owned", "format": "xlsx"})
    assert r_turn_403.status_code == 403

    # 2. CMD creates an export -> 200; E2_HR trying to GET /api/export/{CMD's export_id} -> 403 (BEH-04.50)
    r_ok = cmd_client.post("/api/export", json={"turn_id": "turn_cmd_owned", "format": "xlsx"})
    assert r_ok.status_code == 200
    exp_id = r_ok.json()["export_id"]

    assert e2_client.get(f"/api/export/{exp_id}").status_code == 403
    assert e2_client.get(f"/api/download/{exp_id}").status_code == 403
    assert cmd_client.get(f"/api/export/{exp_id}").status_code == 200
    assert cmd_client.get(f"/api/download/{exp_id}").status_code == 200

    # 3. 24-hour expiry: expire the artifact record and verify GET returns 404
    rec = get_artifact_record(exp_id)
    assert rec is not None
    assert rec.expires_epoch - rec.created_epoch == pytest.approx(24 * 3600)
    rec.expires_epoch = time.time() - 10
    assert cmd_client.get(f"/api/export/{exp_id}").status_code == 404


def test_honest_sheets_and_docs_toast_and_zero_fake_google_urls(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from app.main import app
    import app.tools.export_tool as et

    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(tmp_path / "export_audit.jsonl"))
    client = TestClient(app, headers=_auth("CMD"))

    # Sheets -> downloads xlsx with SHEETS_TOAST, never docs.google.com
    r_sheets = client.post(
        "/api/export",
        json={"format": "sheets", "title": "P&L Statement", "columns": ["fy", "rev"], "rows": [{"fy": "FY26", "rev": 132508}]},
    )
    assert r_sheets.status_code == 200
    j_sheets = r_sheets.json()
    assert j_sheets["format"] == "xlsx"
    assert j_sheets["toast"] == SHEETS_TOAST
    assert "docs.google.com" not in json.dumps(j_sheets)
    assert j_sheets["url"].startswith("/api/export/exp_")

    # Docs -> downloads docx with DOCS_TOAST, never docs.google.com
    r_docs = client.post(
        "/api/export",
        json={"format": "docs", "title": "Board Note", "content": "# Heading\nBody text"},
    )
    assert r_docs.status_code == 200
    j_docs = r_docs.json()
    assert j_docs["format"] == "docx"
    assert j_docs["toast"] == DOCS_TOAST
    assert "docs.google.com" not in json.dumps(j_docs)
    assert j_docs["url"].startswith("/api/export/exp_")

    # Verify export_google_docs is deleted and no docs.google.com URL exists in export_tool.py
    assert not hasattr(et.DocumentExporter, "export_google_docs")
    src = Path(et.__file__).read_text(encoding="utf-8")
    assert "https://docs.google.com" not in src


def test_client_png_svg_audit_endpoint_and_naming_regex_all_formats(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from app.main import app

    audit_file = tmp_path / "export_audit.jsonl"
    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(audit_file))
    client = TestClient(app, headers=_auth("CMD"))

    for fmt in ("png", "svg"):
        r = client.post("/api/export/audit", json={"turn_id": "turn_waterfall_1", "format": fmt, "rows": 10})
        assert r.status_code == 200
        j = r.json()
        assert NAMING_REGEX.fullmatch(j["filename"]), j["filename"]

    rows = read_export_audit_rows(audit_file)
    assert len(rows) == 2
    assert [r["format"] for r in rows] == ["png", "svg"]


def test_r2_m1_client_supplied_rbac_probes_and_no_contract_field(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """M1: Client-supplied rows/content/contract cannot bypass RBAC; ExportRequest rejects contract."""
    from app.main import app, ExportRequest

    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(tmp_path / "export_audit.jsonl"))
    # Verify contract field is deleted from ExportRequest
    assert "contract" not in ExportRequest.model_fields

    e2_client = TestClient(app, headers=_auth("E2_HR"))
    fin_client = TestClient(app, headers=_auth("DIR_FIN"))
    hr_client = TestClient(app, headers=_auth("DIR_HR"))

    # Probe 1: E2_HR client-supplied rows with unrecognized / personal columns -> 403
    r_p1 = e2_client.post(
        "/api/export",
        json={"format": "xlsx", "title": "Leak", "rows": [{"full_name": "X", "staff_no": "99", "secret": "123"}]},
    )
    assert r_p1.status_code == 403

    # Probe 1b: E2_HR client-supplied contract with named statement -> ignored because contract is removed, and empty export -> 400
    r_p1b = e2_client.post(
        "/api/export",
        json={
            "format": "xlsx",
            "contract": {
                "meta": {"domain": "hr"},
                "visual": {
                    "statement": {
                        "meta": {"title": "Salary Slip"},
                        "lines": [{"code": "S1", "label": "Rahul basic_pay", "values": {"FY26": 95000}}],
                    }
                },
            },
        },
    )
    assert r_p1b.status_code in (400, 403, 422)

    # Probe 1c: E2_HR unknown turn_id + client rows fallback -> 403 (no fallback to client rows when turn_id given)
    r_p1c = e2_client.post(
        "/api/export",
        json={"turn_id": "unknown_turn_xyz", "format": "xlsx", "rows": [{"grade": "E1", "headcount": 12}]},
    )
    assert r_p1c.status_code == 403

    # Probe 1d: E2_HR free markdown content export without turn_id -> 403
    r_p1d = e2_client.post(
        "/api/export",
        json={"format": "pdf", "title": "Exfil", "content": "Named roster: Rahul ₹95,000"},
    )
    assert r_p1d.status_code == 403

    # E2_HR client-supplied rows WITH verifiable grade key: personal columns stripped, collapsed to grade/headcount + <5 suppressed
    r_e2_agg = e2_client.post(
        "/api/export",
        json={
            "format": "csv",
            "title": "Grade Distribution",
            "rows": [
                {"grade": "E1", "full_name": "A", "basic_pay": 60000},
                {"grade": "E1", "full_name": "B", "basic_pay": 60000},
                {"grade": "E1", "full_name": "C", "basic_pay": 60000},
                {"grade": "E1", "full_name": "D", "basic_pay": 60000},
                {"grade": "E1", "full_name": "E", "basic_pay": 60000},
                {"grade": "E9", "full_name": "CMD", "basic_pay": 300000},
            ],
        },
    )
    assert r_e2_agg.status_code == 200
    csv_body = e2_client.get(r_e2_agg.json()["url"]).content.decode("utf-8-sig")
    assert "full_name" not in csv_body
    assert "basic_pay" not in csv_body
    assert "60000" not in csv_body
    assert "E9,<5" in csv_body

    # DIR_FIN client-supplied HR compensation rows without grade key -> 403; with grade key -> collapsed to grade/headcount
    r_fin_comp_403 = fin_client.post(
        "/api/export",
        json={"format": "xlsx", "title": "Pay", "rows": [{"emp_id": "ONGC-1", "basic_pay": 60000}]},
    )
    assert r_fin_comp_403.status_code == 403

    r_fin_comp_agg = fin_client.post(
        "/api/export",
        json={"format": "csv", "title": "Pay", "rows": [{"grade": "E1", "basic_pay": 60000}]},
    )
    assert r_fin_comp_agg.status_code == 200
    fin_csv_body = fin_client.get(r_fin_comp_agg.json()["url"]).content.decode("utf-8-sig")
    assert "basic_pay" not in fin_csv_body
    assert "60000" not in fin_csv_body

    # DIR_HR client-supplied named HR rows -> 200
    r_hr_ok = hr_client.post(
        "/api/export",
        json={"format": "csv", "title": "HR Roster", "rows": [{"emp_id": "ONGC-1", "grade": "E1", "basic_pay": 60000}]},
    )
    assert r_hr_ok.status_code == 200


def test_r2_m2_and_m5_fail_closed_download_and_24h_sidecar_sweep(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """M2: Orphan files without sidecars return 404 and are unlinked. M5: 24h TTL honoured + sidecars swept."""
    import os
    from app.main import app, EXPORTS_DIR, EXPORT_TTL_SECONDS, _sweep_old_exports
    from app.export_audit import DEFAULT_EXPORT_TTL_SECONDS, _ARTIFACT_STORE

    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(tmp_path / "export_audit.jsonl"))
    assert DEFAULT_EXPORT_TTL_SECONDS == 86400
    assert EXPORT_TTL_SECONDS == 86400

    # Probe 3b: Orphan file on disk with valid 32-hex id but no artifact record or sidecar -> 404 and unlinked
    orphan_id = "exp_" + "b" * 32
    orphan_file = EXPORTS_DIR / f"{orphan_id}.xlsx"
    orphan_file.write_bytes(b"orphan-bytes")
    _ARTIFACT_STORE.pop(orphan_id, None)

    e2_client = TestClient(app, headers=_auth("E2_HR"))
    r_orphan = e2_client.get(f"/api/export/{orphan_id}")
    assert r_orphan.status_code == 404
    assert not orphan_file.exists()

    # M5: File created 2 hours ago WITH valid 24h sidecar survives _sweep_old_exports()
    cmd_client = TestClient(app, headers=_auth("CMD"))
    r_create = cmd_client.post(
        "/api/export",
        json={"format": "csv", "title": "Retention Check", "rows": [{"fy": "FY26", "rev": 100}]},
    )
    assert r_create.status_code == 200
    exp_id = r_create.json()["export_id"]
    data_path = EXPORTS_DIR / f"{exp_id}.csv"
    sidecar_path = EXPORTS_DIR / f".{exp_id}.meta.json"
    assert data_path.exists() and sidecar_path.exists()

    # Backdate mtime to 2 hours ago, keep 24h expires_epoch -> must NOT be swept
    two_hours_ago = time.time() - 7200
    os.utime(data_path, (two_hours_ago, two_hours_ago))
    _sweep_old_exports()
    assert data_path.exists() and sidecar_path.exists()

    # Expire the sidecar -> both data file and sidecar are swept
    meta = json.loads(sidecar_path.read_text(encoding="utf-8"))
    meta["expires_epoch"] = time.time() - 10
    sidecar_path.write_text(json.dumps(meta), encoding="utf-8")
    _ARTIFACT_STORE.pop(exp_id, None)
    _sweep_old_exports()
    assert not data_path.exists()
    assert not sidecar_path.exists()


def test_r2_m6_session_manager_fallback_ownership_and_full_rows(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """M6: SessionManager fallback checks emp_id + persona and restores full last_rows (>50 rows)."""
    from app.main import app
    from app.session_memory import get_session_manager
    from app.export_audit import _TURN_STORE

    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(tmp_path / "export_audit.jsonl"))
    sm = get_session_manager()
    sid = "sess_m6_test"
    sess = sm.get_or_create(sid, "CMD")
    sess.emp_id = "DEMO-CMD"

    full_75_rows = [{"fy": f"Row_{i}", "pat_cr": i} for i in range(75)]
    sess.last_rows = full_75_rows
    sess.last_columns = ["fy", "pat_cr"]
    capped_contract = {
        "turn_id": "turn_m6_75rows",
        "meta": {"domain": "finance", "as_of": "FY26"},
        "kpis": [],
        "visual": {
            "title": "Seventy Five Rows",
            "table": {
                "columns": ["fy", "pat_cr"],
                "headers": {"fy": "Row", "pat_cr": "PAT"},
                "rows": full_75_rows[:50],
            },
        },
    }
    sess.record_turn(capped_contract, user_query="show 75 rows", all_rows=full_75_rows)
    _TURN_STORE.pop("turn_m6_75rows", None)

    # E2_HR trying to resolve CMD's session turn -> 403
    e2_client = TestClient(app, headers=_auth("E2_HR"))
    assert e2_client.post("/api/export", json={"turn_id": "turn_m6_75rows", "format": "csv"}).status_code == 403

    # CMD resolves via SessionManager fallback and gets all 75 rows
    cmd_client = TestClient(app, headers=_auth("CMD"))
    r_cmd = cmd_client.post("/api/export", json={"turn_id": "turn_m6_75rows", "format": "csv"})
    assert r_cmd.status_code == 200
    assert r_cmd.json()["rows"] == 75


def test_r2_s6_s7_s8_naming_zip_csv_and_audit_ownership(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """S6: Yukti_Export ignored for subject slug. S7: Multi-table CSV returns .zip. S8: /api/export/audit verifies turn ownership."""
    import zipfile
    import io
    from app.main import app

    monkeypatch.setenv("YUKTI_EXPORT_AUDIT_LOG", str(tmp_path / "export_audit.jsonl"))
    register_turn_contract(
        turn_id="turn_s7_multi",
        contract={
            "turn_id": "turn_s7_multi",
            "meta": {"domain": "finance", "as_of": "FY26"},
            "kpis": [],
            "visual": {
                "title": "Yukti_Export",
                "chart": {"title": "Capex Breakdown", "kind": "bar"},
                "tables": [
                    {"title": "Upstream Capex", "columns": ["fy", "capex"], "rows": [{"fy": "FY26", "capex": 100}]},
                    {"title": "Downstream Capex", "columns": ["fy", "capex"], "rows": [{"fy": "FY26", "capex": 50}]},
                ],
            },
        },
        emp_id="DEMO-CMD",
        persona="CMD",
    )

    cmd_client = TestClient(app, headers=_auth("CMD"))
    e2_client = TestClient(app, headers=_auth("E2_HR"))

    # S8: E2_HR trying to audit CMD's turn_id -> 403
    assert e2_client.post("/api/export/audit", json={"turn_id": "turn_s7_multi", "format": "png"}).status_code == 403

    # S6 + S7: CMD exporting multi-table turn as CSV gets a .zip with capex_breakdown slug (not yukti_export)
    r_zip = cmd_client.post("/api/export", json={"turn_id": "turn_s7_multi", "format": "csv"})
    assert r_zip.status_code == 200
    j_zip = r_zip.json()
    assert j_zip["filename"].endswith(".zip")
    assert "capex_breakdown" in j_zip["filename"].lower()
    assert NAMING_REGEX.fullmatch(j_zip["filename"]), j_zip["filename"]

    zip_bytes = cmd_client.get(j_zip["url"]).content
    with zipfile.ZipFile(io.BytesIO(zip_bytes), "r") as zf:
        names = zf.namelist()
        assert len(names) == 2
        assert all(n.endswith(".csv") for n in names)

