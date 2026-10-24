import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from typing import Dict, Any

def write_dcf_xlsx(result: Dict[str, Any], filename: str):
    wb = openpyxl.Workbook()
    
    # 1. Cover
    ws_cover = wb.active
    ws_cover.title = "Cover"
    ws_cover["A1"] = "ONGC DCF Valuation"
    ws_cover["A1"].font = Font(bold=True, size=14)
    ws_cover["A3"] = result.get("disclaimer", "Illustrative valuation, not investment advice.")
    ws_cover["A3"].font = Font(bold=True)
    ws_cover["A5"] = f"Run ID: {result.get('run_id')}"
    ws_cover["A6"] = f"As Of: {result.get('as_of')}"
    ws_cover["A8"] = "Provenance: Illustrative"

    # 2. Assumptions
    ws_a = wb.create_sheet("Assumptions")
    ws_a["A1"] = "Name"
    ws_a["B1"] = "Value"
    ws_a["C1"] = "Source"
    
    yellow = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    
    a_dict = result.get("assumptions", {})
    row = 2
    for k, v in a_dict.items():
        if k == "forecast_fcff":
            continue
        if isinstance(v, dict) and "value" in v:
            val = v["value"]
            src = v["source"]
        else:
            val = v
            src = "Engine"
            
        ws_a[f"A{row}"] = k
        ws_a[f"B{row}"] = val
        if k in ["wacc_pct", "terminal_g_pct", "brent_usd", "capex_scale_pct", "production_decline_pct", "holdco_discount_pct", "exit_ev_ebitda"]:
            ws_a[f"B{row}"].fill = yellow
        ws_a[f"C{row}"] = src
        # Create named range
        # Use Workbook.defined_names
        from openpyxl.workbook.defined_name import DefinedName
        dn = DefinedName(k, attr_text=f"Assumptions!$B${row}")
        wb.defined_names.add(dn)
        row += 1

    # 3. WACC
    ws_wacc = wb.create_sheet("WACC")
    ws_wacc["A1"] = "WACC Build"
    w_row = 2
    for k, v in result.get("wacc_build", {}).items():
        ws_wacc[f"A{w_row}"] = k
        ws_wacc[f"B{w_row}"] = v
        w_row += 1

    # 4. Historical
    ws_hist = wb.create_sheet("Historical")
    ws_hist.append(["fy", "line_code", "value_inr_cr", "source_doc", "page_or_source"])
    
    for row_data in result.get("historical", []):
        ws_hist.append([
            row_data.get("fy", ""),
            row_data.get("line_code", ""),
            row_data.get("value_inr_cr", ""),
            row_data.get("source_doc", ""),
            row_data.get("page", row_data.get("source", ""))
        ])

    # 5. Forecast
    ws_f = wb.create_sheet("Forecast")
    ws_f["A1"] = "Metric"
    ws_f["A2"] = "t"
    ws_f["A3"] = "FCFF"
    
    fcff_data = result.get("forecast", [])
    
    if isinstance(fcff_data, dict):
        years = fcff_data.get("years", [])
        lines = fcff_data.get("lines", {})
        fcff_list = lines.get("fcff", [])
        if not fcff_list and "rows" in fcff_data:
            # Maybe it's in rows
            fcff_row = next((r for r in fcff_data["rows"] if r.get("code") == "fcff"), {})
            fcff_list = fcff_row.get("values", [])
            
        for i, val in enumerate(fcff_list):
            col = chr(ord('B') + i)
            period = years[i] if i < len(years) else f"FY{27+i}"
            ws_f[f"{col}1"] = period
            ws_f[f"{col}2"] = i + 1
            ws_f[f"{col}3"] = val
        list_len = list_len
    else:
        for i, r in enumerate(fcff_data):
            col = chr(ord('B') + i)
            ws_f[f"{col}1"] = r.get("period", f"FY{27+i}")
            ws_f[f"{col}2"] = r.get("t", i+1)
            ws_f[f"{col}3"] = r.get("fcff", 0)
        list_len = len(fcff_data)

    last_col = chr(ord('B') + list_len - 1) if list_len > 0 else 'B'

    # 6. Valuation
    ws_v = wb.create_sheet("Valuation")
    
    ws_v["A1"] = "Sum of PV FCFF"
    # SUMPRODUCT(Forecast!B3:last3 / (1+wacc_pct/100)^Forecast!B2:last2)
    ws_v["B1"] = f"=SUMPRODUCT(Forecast!B3:{last_col}3 / (1+wacc_pct/100)^Forecast!B2:{last_col}2)"
    
    ws_v["A2"] = "TV (Gordon)"
    # FCFF_last * (1+g) / (wacc - g)
    ws_v["B2"] = f"=Forecast!{last_col}3 * (1+terminal_g_pct/100) / (wacc_pct/100 - terminal_g_pct/100)"
    
    ws_v["A3"] = "PV of TV"
    # TV / (1+wacc)^10
    ws_v["B3"] = f"=B2 / (1+wacc_pct/100)^{list_len}"
    
    ws_v["A4"] = "EV"
    ws_v["B4"] = "=B1 + B3"
    
    ws_v["A5"] = "Less: Borrowings"
    ws_v["B5"] = "=borrowings_inr_cr"
    
    ws_v["A6"] = "Less: Lease Liab"
    ws_v["B6"] = "=lease_liab_inr_cr"
    
    ws_v["A7"] = "Less: Decom Prov"
    ws_v["B7"] = "=decom_provision_inr_cr"
    
    ws_v["A8"] = "Add: Cash & CCE"
    ws_v["B8"] = "=cash_and_equivalents_inr_cr"
    
    ws_v["A9"] = "Add: SRF Deposits"
    ws_v["B9"] = "=srf_deposits_inr_cr"
    
    ws_v["A10"] = "Add: Listed Inv"
    ws_v["B10"] = "=listed_investments_inr_cr * (1 - holdco_discount_pct/100)"
    
    ws_v["A11"] = "Add: Unlisted Inv"
    ws_v["B11"] = "=unlisted_investments_inr_cr"
    
    ws_v["A12"] = "Equity Before Discount"
    ws_v["B12"] = "=B4 - B5 - B6 - B7 + B8 + B9 + B10 + B11"
    
    ws_v["A13"] = "Equity Value"
    ws_v["B13"] = "=B12"
    
    ws_v["A14"] = "Shares Outstanding"
    ws_v["B14"] = "=shares_outstanding"
    
    ws_v["A15"] = "Per Share Value"
    ws_v["B15"] = "=B13 / B14"
    dn_ps = DefinedName("PerShareValue", attr_text="Valuation!$B$15")
    wb.defined_names.add(dn_ps)

    # 7. Sensitivity
    ws_s = wb.create_sheet("Sensitivity")
    ws_s["A1"] = "WACC \\ g"
    
    # Need to hardcode the axis base values for the formulas to be purely excel-based,
    # or reference wacc_pct and terminal_g_pct. The spec asks for a 9x9 grid. 
    # Let's put the axis values directly as numbers or formulas from base.
    # WACC axis in A3:A11, g axis in B2:J2
    w_base = a_dict.get("wacc_pct", {}).get("value", 10.0)
    g_base = a_dict.get("terminal_g_pct", {}).get("value", 3.0)
    
    for i in range(-4, 5):
        row = i + 7 # WACC at row 3 to 11
        ws_s.cell(row=row, column=1, value=f"=wacc_pct + {i*0.5}")
        
    for j in range(-4, 5):
        col = j + 6 # g at col 2 (B) to 10 (J)
        ws_s.cell(row=2, column=col, value=f"=terminal_g_pct + {j*0.5}")
        
    for r in range(3, 12):
        for c in range(2, 11):
            w_ref = f"$A{r}"
            g_ref = f"{openpyxl.utils.get_column_letter(c)}$2"
            
            # (SUMPRODUCT(FCFF/(1+W)^t) + TV/(1+W)^t + adjustments) / shares
            # using named ranges for adjustments
            sumpv = f"SUMPRODUCT(Forecast!$B$3:${last_col}$3 / (1+{w_ref}/100)^Forecast!$B$2:${last_col}$2)"
            tv = f"(Forecast!${last_col}$3 * (1+{g_ref}/100) / ({w_ref}/100 - {g_ref}/100)) / (1+{w_ref}/100)^{list_len}"
            adj = "-borrowings_inr_cr - lease_liab_inr_cr - decom_provision_inr_cr + cash_and_equivalents_inr_cr + srf_deposits_inr_cr + listed_investments_inr_cr * (1 - holdco_discount_pct/100) + unlisted_investments_inr_cr"
            eq = f"({sumpv} + {tv} {adj})"
            ps = f"={eq} / shares_outstanding"
            ws_s.cell(row=r, column=c, value=ps)

    # 8. Tornado
    ws_t = wb.create_sheet("Tornado")
    ws_t["A1"] = f"Computed by engine run {result.get('run_id')}"
    ws_t["A2"] = "Driver"
    ws_t["B2"] = "Low (INR)"
    ws_t["C2"] = "High (INR)"
    
    t_row = 3
    for tor in result.get("tornado", []):
        ws_t[f"A{t_row}"] = tor.get("driver")
        ws_t[f"B{t_row}"] = tor.get("low_inr")
        ws_t[f"C{t_row}"] = tor.get("high_inr")
        t_row += 1
        
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Sensitivity Tornado"
    data = Reference(ws_t, min_col=2, min_row=2, max_row=t_row-1, max_col=3)
    cats = Reference(ws_t, min_col=1, min_row=3, max_row=t_row-1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws_t.add_chart(chart, "E2")
    
    # 9. Sources
    ws_src = wb.create_sheet("Sources")
    ws_src["A1"] = "Lineage / Item"
    ws_src["B1"] = "Source"
    s_row = 2
    
    for ref in result.get("lineage_refs", []):
        ws_src[f"A{s_row}"] = ref.get("item", ref.get("text", ""))
        ws_src[f"B{s_row}"] = ref.get("source", ref.get("doc_id", ""))
        s_row += 1
        
    for br in result.get("bridge", []):
        ws_src[f"A{s_row}"] = br.get("item", br.get("label", ""))
        ws_src[f"B{s_row}"] = br.get("source", "")
        s_row += 1
        
    for hist in result.get("historical", []):
        ws_src[f"A{s_row}"] = f"{hist.get('fy', '')} {hist.get('line_code', '')}"
        ws_src[f"B{s_row}"] = hist.get("source_doc", hist.get("source", ""))
        s_row += 1
    wb.save(filename)
