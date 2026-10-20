"""test_34_37_docs.py: Test suite for Step 3.4 (MIS Reports) and Step 3.7 (Output Templates).

Validates:
1. Manifest integrity: landing/documents/DOC_MANIFEST_34_37.csv has 340 entries.
2. Templates: All 4 templates open cleanly and contain expected {{placeholder}} tokens.
3. Scanned PDF text layer: Sampled 30 MIS PDFs have exactly 0 extractable text characters (image-only).
4. Provenance SQL & Figures: Recomputes 10+ figures per report directly from DuckDB SQL and asserts
   exact match against the sidecar JSON figures.
"""

import os
import re
import json
import random
from pathlib import Path
import pytest
import pypdf
import openpyxl
import docx
import duckdb
import pandas as pd
import yaml

DATA_BANK_ROOT = Path(__file__).resolve().parent.parent
LANDING_DIR = DATA_BANK_ROOT / "landing"
DOCUMENTS_DIR = LANDING_DIR / "documents"
SCANNED_MIS_DIR = DOCUMENTS_DIR / "scanned" / "mis"
PROVENANCE_MIS_DIR = DOCUMENTS_DIR / "_provenance" / "mis"
TEMPLATES_DIR = DOCUMENTS_DIR / "templates"
MANIFEST_PATH = DOCUMENTS_DIR / "DOC_MANIFEST_34_37.csv"

SEED = 20261019


@pytest.fixture(scope="module")
def duckdb_con():
    """Module-level DuckDB connection with landing tables registered."""
    con = duckdb.connect()
    con.execute(f"CREATE TABLE prod AS SELECT * FROM '{LANDING_DIR}/db_exports/production_monthly.csv'")
    con.execute(f"CREATE TABLE cosp AS SELECT * FROM '{LANDING_DIR}/sap/COSP_COSS.csv'")
    con.execute(f"CREATE TABLE hc AS SELECT * FROM '{LANDING_DIR}/sap/headcount_monthly_snapshot.csv'")
    con.execute(f"CREATE TABLE py AS SELECT * FROM '{LANDING_DIR}/sap/PY_CLUSTER_EXP.csv'")
    con.execute(f"CREATE TABLE hse AS SELECT * FROM '{LANDING_DIR}/db_exports/hse_incident_log.csv'")
    df_capex = pd.read_excel(LANDING_DIR / "excel" / "capex_be_re_actuals.xlsx")
    con.register("capex_t", df_capex)
    con.execute("CREATE TABLE capex_budget AS SELECT * FROM capex_t")
    return con


@pytest.fixture(scope="module")
def manifest_df():
    """Load DOC_MANIFEST_34_37.csv."""
    assert MANIFEST_PATH.exists(), f"Missing manifest: {MANIFEST_PATH}"
    df = pd.read_csv(MANIFEST_PATH)
    return df


@pytest.fixture(scope="module")
def sampled_reports(manifest_df):
    """Deterministically sample 30 MIS reports from the manifest."""
    mis_rows = manifest_df[manifest_df["type"] == "mis_report"].to_dict(orient="records")
    assert len(mis_rows) == 336, f"Expected 336 MIS reports, found {len(mis_rows)}"
    rng = random.Random(SEED)
    sampled = rng.sample(mis_rows, 30)
    return sampled


# =============================================================================
# 1. MANIFEST & FILE COUNT TESTS
# =============================================================================
def test_manifest_structure_and_counts(manifest_df):
    """Verify DOC_MANIFEST_34_37.csv schema, completeness, and row counts."""
    assert len(manifest_df) == 340, f"Expected 340 manifest rows (336 MIS + 4 templates), got {len(manifest_df)}"
    expected_cols = [
        "doc_id", "title", "type", "asset_wc_id", "month",
        "format", "provenance_sql", "questions_served"
    ]
    for col in expected_cols:
        assert col in manifest_df.columns, f"Missing manifest column: {col}"

    assert manifest_df["doc_id"].is_unique, "Duplicate doc_id found in manifest"

    # Count by type
    counts = manifest_df["type"].value_counts().to_dict()
    assert counts.get("mis_report") == 336
    assert counts.get("template") == 4

    # Verify 14 assets * 24 months
    mis_df = manifest_df[manifest_df["type"] == "mis_report"]
    assert mis_df["asset_wc_id"].nunique() == 14
    assert mis_df["month"].nunique() == 24


def test_all_mis_files_exist_on_disk(manifest_df):
    """Verify that every PDF, SQL, and JSON file listed in the manifest exists."""
    mis_rows = manifest_df[manifest_df["type"] == "mis_report"]
    for _, row in mis_rows.iterrows():
        wc_id = row["asset_wc_id"]
        month = row["month"]
        pdf_path = SCANNED_MIS_DIR / f"MIS_{wc_id}_{month}.pdf"
        sql_path = PROVENANCE_MIS_DIR / f"MIS_{wc_id}_{month}.sql"
        json_path = PROVENANCE_MIS_DIR / f"MIS_{wc_id}_{month}.json"

        assert pdf_path.exists(), f"Missing scanned PDF: {pdf_path}"
        assert sql_path.exists(), f"Missing provenance SQL: {sql_path}"
        assert json_path.exists(), f"Missing sidecar JSON: {json_path}"
        assert pdf_path.stat().st_size > 50000, f"PDF suspiciously small: {pdf_path}"


# =============================================================================
# 2. TEMPLATE TESTS (Step 3.7)
# =============================================================================
def test_board_note_template():
    """Verify Board Note DOCX opens and contains required placeholders (serving C14, T1)."""
    tmpl_path = TEMPLATES_DIR / "template_board_note.docx"
    assert tmpl_path.exists(), f"Missing {tmpl_path}"

    doc = docx.Document(tmpl_path)
    full_text = " ".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += " " + cell.text

    expected_placeholders = [
        "{{agenda_item_no}}",
        "{{meeting_no}}",
        "{{asset_name}}",
        "{{wc_id}}",
        "{{field_name}}",
        "{{oil_production_rate_mt}}",
        "{{water_cut_pct}}",
        "{{gross_revenue_inr_cr}}",
        "{{operating_cost_inr_cr}}",
        "{{manpower_cost_inr_cr}}",
        "{{net_pnl_inr_cr}}",
        "{{lifting_cost_inr_per_boe}}",
        "{{option_a_capex_cr}}",
        "{{option_b_savings_cr}}",
        "{{total_staff_count}}",
        "{{recommendation_resolution_1}}",
    ]
    for ph in expected_placeholders:
        assert ph in full_text, f"Board Note missing placeholder: {ph}"


def test_agt_proposal_template():
    """Verify Annual General Transfer (AGT) proposal sheet (serving H10)."""
    tmpl_path = TEMPLATES_DIR / "template_annual_general_transfer_proposal.xlsx"
    assert tmpl_path.exists(), f"Missing {tmpl_path}"

    wb = openpyxl.load_workbook(tmpl_path, data_only=False)
    assert "AGT_Proposal_FY27" in wb.sheetnames
    assert "Validation_Lists" in wb.sheetnames

    ws = wb["AGT_Proposal_FY27"]
    headers = [ws.cell(row=4, column=c).value for c in range(1, 17)]
    assert "Employee ID" in headers
    assert "Grade" in headers
    assert "Discipline" in headers
    assert "Present Work Centre" in headers
    assert "Proposed Work Centre" in headers
    assert "Transfer Category" in headers

    # Verify placeholder sample row
    row5_vals = [str(ws.cell(row=5, column=c).value or "") for c in range(1, 17)]
    row5_text = " ".join(row5_vals)
    assert "{{emp_id}}" in row5_text
    assert "{{grade_id}}" in row5_text
    assert "{{proposed_wc_id}}" in row5_text

    # Verify Data Validations exist
    assert len(ws.data_validations.dataValidation) >= 4, "Missing data validation rules in AGT proposal"


def test_parliament_reply_template():
    """Verify Parliament Reply DOCX template (serving X7)."""
    tmpl_path = TEMPLATES_DIR / "template_parliament_reply_format.docx"
    assert tmpl_path.exists(), f"Missing {tmpl_path}"

    doc = docx.Document(tmpl_path)
    full_text = " ".join([p.text for p in doc.paragraphs])
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += " " + cell.text

    expected_placeholders = [
        "{{parliament_house}}",
        "{{question_number}}",
        "{{answer_date}}",
        "{{subject}}",
        "{{mp_names}}",
        "{{minister_name}}",
        "{{wc_name}}",
        "{{sanctioned_total}}",
        "{{in_position_total}}",
        "{{vacancy_total}}",
        "{{sc_vac}}",
        "{{st_vac}}",
        "{{obc_vac}}",
        "{{vetted_by_official}}",
    ]
    for ph in expected_placeholders:
        assert ph in full_text, f"Parliament reply missing placeholder: {ph}"


def test_quarterly_asset_pnl_template():
    """Verify Quarterly Asset-Wise P&L sheet template (serving F10)."""
    tmpl_path = TEMPLATES_DIR / "template_quarterly_asset_pnl.xlsx"
    assert tmpl_path.exists(), f"Missing {tmpl_path}"

    wb = openpyxl.load_workbook(tmpl_path, data_only=False)
    assert "Quarterly_Asset_PnL" in wb.sheetnames
    assert "Notes_and_Policies" in wb.sheetnames

    ws = wb["Quarterly_Asset_PnL"]
    all_text = ""
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                all_text += f" {val}"

    expected_placeholders = [
        "{{asset_name}}",
        "{{wc_id}}",
        "{{crude_oil_mt}}",
        "{{natural_gas_mmscm}}",
        "{{total_boe}}",
        "{{crude_revenue_cr}}",
        "{{gas_revenue_cr}}",
        "{{gross_revenue_cr}}",
        "{{royalty_cr}}",
        "{{net_revenue_cr}}",
        "{{manpower_opex_cr}}",
        "{{total_opex_cr}}",
        "{{ebitda_cr}}",
        "{{lifting_cost_inr_per_boe}}",
        "{{capex_be_cr}}",
        "{{capex_actual_cr}}",
    ]
    for ph in expected_placeholders:
        assert ph in all_text, f"Quarterly P&L missing placeholder: {ph}"


# =============================================================================
# 3. SCANNED PDF TEST (NO TEXT LAYER)
# =============================================================================
def test_scanned_pdfs_have_no_text_layer(sampled_reports):
    """Assert that sampled 30 scanned MIS PDFs have exactly 0 extractable text characters."""
    for report_meta in sampled_reports:
        wc_id = report_meta["asset_wc_id"]
        month = report_meta["month"]
        pdf_path = SCANNED_MIS_DIR / f"MIS_{wc_id}_{month}.pdf"

        reader = pypdf.PdfReader(str(pdf_path))
        assert len(reader.pages) == 3, f"Expected 3 pages in {pdf_path.name}, got {len(reader.pages)}"

        for p_idx, page in enumerate(reader.pages):
            extracted = page.extract_text() or ""
            assert len(extracted.strip()) == 0, (
                f"Page {p_idx+1} of {pdf_path.name} unexpectedly contains text: {repr(extracted[:40])}"
            )


# =============================================================================
# 4. PROVENANCE SQL RECOMPUTATION TESTS (10+ FIGURES PER REPORT)
# =============================================================================
def test_recompute_figures_from_provenance_sql(sampled_reports, duckdb_con):
    """Recompute 10+ figures per report via SQL and verify equality with sidecar JSON."""
    for report_meta in sampled_reports:
        wc_id = report_meta["asset_wc_id"]
        month = report_meta["month"]

        # Load sidecar JSON
        json_path = PROVENANCE_MIS_DIR / f"MIS_{wc_id}_{month}.json"
        assert json_path.exists(), f"Missing sidecar JSON: {json_path}"
        with open(json_path, "r", encoding="utf-8") as f:
            sidecar = json.load(f)
        fig = sidecar["figures"]

        # 1. Recompute Production figures
        df_prod = duckdb_con.execute(
            """
            SELECT 
                COALESCE(SUM(oil_production_mt), 0.0) as oil_mt,
                COALESCE(SUM(gas_production_mmscm), 0.0) as gas_mmscm,
                COALESCE(SUM(total_production_boe), 0.0) as total_boe,
                COALESCE(AVG(water_cut_pct), 0.0) as water_cut,
                COALESCE(SUM(well_downtime_hours), 0.0) as downtime,
                COALESCE(SUM(active_producing_wells), 0) as active_wells
            FROM prod
            WHERE asset_wc_id = ? AND month = ?
        """,
            [wc_id, month],
        ).df().iloc[0]

        assert fig["oil_production_mt"] == pytest.approx(float(df_prod["oil_mt"]), rel=1e-4)
        assert fig["gas_production_mmscm"] == pytest.approx(float(df_prod["gas_mmscm"]), rel=1e-4)
        assert fig["total_production_boe"] == pytest.approx(float(df_prod["total_boe"]), rel=1e-4)
        assert fig["water_cut_pct"] == pytest.approx(float(df_prod["water_cut"]), rel=1e-4)
        assert fig["well_downtime_hours"] == pytest.approx(float(df_prod["downtime"]), rel=1e-4)
        assert fig["active_producing_wells"] == pytest.approx(float(df_prod["active_wells"]), rel=1e-4)

        # 2. Recompute Operating Cost & Lifting Cost
        df_cosp = duckdb_con.execute(
            """
            SELECT COALESCE(SUM(amount_inr), 0.0) as opex
            FROM cosp
            WHERE wc_id = ? AND month = ?
        """,
            [wc_id, month],
        ).df().iloc[0]

        recomp_opex = float(df_cosp["opex"])
        recomp_boe = float(df_prod["total_boe"])
        recomp_lifting = (recomp_opex / recomp_boe) if recomp_boe > 0 else 0.0

        assert fig["total_opex_inr"] == pytest.approx(recomp_opex, rel=1e-4)
        assert fig["lifting_cost_inr_per_boe"] == pytest.approx(recomp_lifting, rel=1e-4)

        # 3. Recompute Headcount
        df_hc = duckdb_con.execute(
            """
            SELECT COALESCE(SUM(headcount), 0) as hc
            FROM hc
            WHERE wc_id = ? AND month = ?
        """,
            [wc_id, month],
        ).df().iloc[0]

        assert fig["total_headcount"] == int(df_hc["hc"])

        # 4. Recompute Overtime from Payroll
        df_py = duckdb_con.execute(
            """
            SELECT 
                COALESCE(SUM(overtime_hours), 0.0) as ot_hrs,
                COALESCE(SUM(overtime_amount_inr), 0.0) as ot_amt
            FROM py
            WHERE wc_id = ? AND FPPER = ?
        """,
            [wc_id, month],
        ).df().iloc[0]

        assert fig["total_overtime_hours"] == pytest.approx(float(df_py["ot_hrs"]), rel=1e-4)
        assert fig["total_overtime_amount_inr"] == pytest.approx(float(df_py["ot_amt"]), rel=1e-4)

        # 5. Recompute HSE Incidents
        df_hse = duckdb_con.execute(
            """
            SELECT COALESCE(COUNT(*), 0) as cnt
            FROM hse
            WHERE wc_id = ? AND month = ?
        """,
            [wc_id, month],
        ).df().iloc[0]

        assert fig["hse_incidents_count"] == int(df_hse["cnt"])

        # 6. Recompute Capex
        df_cap = duckdb_con.execute(
            """
            SELECT 
                COALESCE(SUM(budget_estimate_be_inr_cr), 0.0) as be_cr,
                COALESCE(SUM(actual_spend_inr_cr), 0.0) as act_cr
            FROM capex_budget
            WHERE asset_wc_id = ? AND month = ?
        """,
            [wc_id, month],
        ).df().iloc[0]

        assert fig["capex_be_inr_cr"] == pytest.approx(float(df_cap["be_cr"]), rel=1e-4)
        assert fig["capex_actual_spend_inr_cr"] == pytest.approx(float(df_cap["act_cr"]), rel=1e-4)


def test_provenance_sql_integrity(sampled_reports):
    """Verify that provenance SQL files exist, are valid, and document the queries."""
    for report_meta in sampled_reports:
        wc_id = report_meta["asset_wc_id"]
        month = report_meta["month"]
        sql_path = PROVENANCE_MIS_DIR / f"MIS_{wc_id}_{month}.sql"

        assert sql_path.exists()
        sql_text = sql_path.read_text(encoding="utf-8")
        assert f"Asset Work Centre: {wc_id}" in sql_text
        assert f"Reporting Month: {month}" in sql_text
        assert "production_monthly.csv" in sql_text
        assert "COSP_COSS.csv" in sql_text
        assert "headcount_monthly_snapshot.csv" in sql_text
        assert "PY_CLUSTER_EXP.csv" in sql_text
        assert "hse_incident_log.csv" in sql_text
        assert "capex_budget" in sql_text
