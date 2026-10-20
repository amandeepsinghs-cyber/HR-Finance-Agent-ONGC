"""gen_37_templates.py: Output template generator for Step 3.7.

Generates standard ONGC corporate output templates with placeholders ({{field_name}}):
1. landing/documents/templates/template_board_note.docx (DOCX)
2. landing/documents/templates/template_annual_general_transfer_proposal.xlsx (XLSX)
3. landing/documents/templates/template_parliament_reply_format.docx (DOCX)
4. landing/documents/templates/template_quarterly_asset_pnl.xlsx (XLSX)
"""

import os
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
TEMPLATES_DIR = PROJECT_ROOT / "landing" / "documents" / "templates"


def set_cell_background(cell, hex_color: str):
    """Set background shading for a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def generate_board_note_template(out_path: Path):
    """Generate the Board Note template (DOCX with placeholders)."""
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Header banner
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hdr = p_hdr.add_run("OIL AND NATURAL GAS CORPORATION LIMITED\n")
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(14)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = RGBColor(0, 51, 102)

    r_sub = p_hdr.add_run("BOARD OF DIRECTORS | AGENDA NOTE\nSTRICTLY CONFIDENTIAL - FOR BOARD CONSIDERATION ONLY")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(180, 0, 0)

    # Metadata table
    tbl = doc.add_table(rows=4, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Agenda Item No.", "{{agenda_item_no}}", "Meeting No. & Date", "{{meeting_no}} dated {{meeting_date}}"),
        ("Sponsoring Directorate", "{{sponsoring_directorate}}", "Target Asset / Field", "{{asset_name}} ({{wc_id}}) / {{field_name}}"),
        ("Subject", "{{subject}}", "Classification", "Confidential / Board Note"),
        ("Prepared By", "{{prepared_by_name}}, {{prepared_by_designation}}", "Vetted By", "{{vetted_by_director}}, Director ({{directorate_name}})"),
    ]

    for i, (k1, v1, k2, v2) in enumerate(meta_rows):
        row = tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(3.4)
        c1.width = Inches(3.4)
        set_cell_background(c0, "F0F4F8")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(2)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(f"{k1}: ")
        r0.font.bold = True
        r0.font.size = Pt(8.5)
        p0.add_run(v1).font.size = Pt(8.5)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(2)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(f"{k2}: ")
        r1.font.bold = True
        r1.font.size = Pt(8.5)
        p1.add_run(v2).font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 1: Executive Summary
    h1 = doc.add_heading("1. Proposal & Executive Summary", level=1)
    h1.style.font.name = "Arial"
    h1.style.font.color.rgb = RGBColor(0, 51, 102)
    p1 = doc.add_paragraph("{{executive_summary}}")
    p1.paragraph_format.line_spacing = 1.15

    # Section 2: Operational & Reservoir Performance
    h2 = doc.add_heading("2. Operational Performance & Reservoir Profile", level=1)
    h2.style.font.name = "Arial"
    h2.style.font.color.rgb = RGBColor(0, 51, 102)
    p2 = doc.add_paragraph(
        "Field {{field_name}} located under {{asset_name}} ({{wc_id}}) has exhibited the following operational trends:\n"
        "• Current Production Rate: {{oil_production_rate_mt}} MT/month Crude Oil; {{gas_production_rate_mmscmd}} MMSCM/month Natural Gas.\n"
        "• Water Cut & Reservoir Pressure: Current water cut stands at {{water_cut_pct}}% with reservoir pressure status noted as {{reservoir_pressure_status}}.\n"
        "• Active Well Count & Downtime: {{active_wells_count}} active producing wells with {{well_downtime_hours}} downtime hours recorded.\n"
        "• Root Cause Analysis: {{root_cause_reservoir_analysis}}"
    )
    p2.paragraph_format.line_spacing = 1.15

    # Section 3: Financial Performance
    h3 = doc.add_heading("3. Financial Analysis & Field Economics", level=1)
    h3.style.font.name = "Arial"
    h3.style.font.color.rgb = RGBColor(0, 51, 102)

    fin_tbl = doc.add_table(rows=7, cols=3)
    fin_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    fin_data = [
        ("Financial Metric", "Value (INR)", "Operational Benchmark / Reference"),
        ("Gross Operating Revenue", "₹{{gross_revenue_inr_cr}} Crores", "Realized hydrocarbon sales net of discounts"),
        ("Statutory Levies (Royalty & Cess)", "₹{{statutory_levies_inr_cr}} Crores", "Statutory obligations to Central/State Govts"),
        ("Operating Expenditure (OPEX)", "₹{{operating_cost_inr_cr}} Crores", "Surface facilities, workover, chemicals, power"),
        ("Allocated Manpower Cost", "₹{{manpower_cost_inr_cr}} Crores", "Direct payroll and terminal benefits"),
        ("Net Operating Profit / (Loss)", "₹{{net_pnl_inr_cr}} Crores", "EBITDA contribution of field operations"),
        ("Lifting Cost per BOE", "₹{{lifting_cost_inr_per_boe}} / BOE", "Asset average: ₹{{asset_avg_lifting_cost_inr}} / BOE"),
    ]
    for r_idx, row in enumerate(fin_data):
        for c_idx, val in enumerate(row):
            cell = fin_tbl.rows[r_idx].cells[c_idx]
            cell.paragraphs[0].text = val
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            if r_idx == 0:
                set_cell_background(cell, "003366")
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "F9FBFD" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 40, 40, 80, 80)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 4: Strategic Turnaround & Divestment Options
    h4 = doc.add_heading("4. Strategic Options Evaluated", level=1)
    h4.style.font.name = "Arial"
    h4.style.font.color.rgb = RGBColor(0, 51, 102)
    p4 = doc.add_paragraph(
        "To mitigate ongoing operational deficits, three alternative options have been evaluated:\n\n"
        "• Option A - In-House Infill Drilling & Secondary Recovery:\n"
        "  Requires capital expenditure of ₹{{option_a_capex_cr}} Crores over 24 months. Expected NPV: ₹{{option_a_npv_cr}} Crores.\n\n"
        "• Option B - Production Enhancement Contract (PEC) / Farm-Out:\n"
        "  Transfer field operations to external specialized operator under revenue-sharing model. Annual OPEX savings: ₹{{option_b_savings_cr}} Crores. Projected 5-year payoff: ₹{{option_b_payoff_cr}} Crores.\n\n"
        "• Option C - Field Divestment / Plug & Abandonment:\n"
        "  Complete decommissioning and plugging of {{active_wells_count}} wells. Estimated one-time abandonment liability: ₹{{option_c_cost_cr}} Crores."
    )
    p4.paragraph_format.line_spacing = 1.15

    # Section 5: Manpower Redeployment
    h5 = doc.add_heading("5. Manpower Transition & Redeployment Plan", level=1)
    h5.style.font.name = "Arial"
    h5.style.font.color.rgb = RGBColor(0, 51, 102)
    p5 = doc.add_paragraph(
        "Currently, {{total_staff_count}} personnel ({{executive_staff_count}} Executives and {{non_executive_staff_count}} Non-Executives) "
        "are deployed at {{field_name}}. In accordance with JRTP guidelines, personnel will be redeployed as follows:\n"
        "• {{redeployment_target_work_centres}}\n"
        "• All redeployments will satisfy mandated Category-C station balance and maintain business continuity across high-performing assets."
    )
    p5.paragraph_format.line_spacing = 1.15

    # Section 6: Recommendation & Board Resolution
    h6 = doc.add_heading("6. Recommendation & Decision Sought", level=1)
    h6.style.font.name = "Arial"
    h6.style.font.color.rgb = RGBColor(0, 51, 102)
    p6 = doc.add_paragraph(
        "The Board is requested to approve:\n"
        "1. {{recommendation_resolution_1}}\n"
        "2. {{recommendation_resolution_2}}\n"
        "3. Authorization to CMD, ONGC to finalize commercial and operational agreements as per BDP guidelines."
    )
    p6.paragraph_format.line_spacing = 1.15

    # Sign-off block
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    sig_p = doc.add_paragraph()
    sig_p.add_run("Prepared & Submitted by: {{prepared_by_name}}, {{prepared_by_designation}}\n").font.bold = True
    sig_p.add_run("Recommended by: {{recommended_by_director}}, Director ({{directorate_name}})\n").font.bold = True
    sig_p.add_run("Approved for Board Consideration: {{approved_by_cmd}}, Chairman & Managing Director\n").font.bold = True

    doc.save(out_path)


def generate_agt_proposal_template(out_path: Path):
    """Generate the Annual General Transfer (AGT) proposal workbook (XLSX with validation lists)."""
    wb = openpyxl.Workbook()

    # Sheet 1: AGT Proposal
    ws1 = wb.active
    ws1.title = "AGT_Proposal_FY27"
    ws1.views.sheetView[0].showGridLines = True

    # Sheet 2: Validation Reference Data
    ws2 = wb.create_sheet(title="Validation_Lists")
    ws2.views.sheetView[0].showGridLines = True

    # Populate validation lists
    grades = ["E0", "E1", "E2", "E3", "E4", "E5", "E6", "E7", "E8", "E9"]
    disciplines = [
        "Geology", "Geophysics", "Reservoir", "Drilling", "Production",
        "Mechanical", "Electrical", "Instrumentation", "Chemical", "Civil",
        "F&A", "HR", "MM", "Legal", "Medical", "Logistics", "Corporate",
    ]
    tenure_statuses = [
        "Overdue (>5 yrs)", "Due for Rotation (3-5 yrs)", "Within Normal Tenure (<3 yrs)",
        "Retirement Exemption (<2 yrs to 60)", "Medical/Hardship Exemption",
    ]
    hard_station_statuses = [
        "Completed Category C/Offshore", "Pending Mandatory Category C",
        "Currently Serving Hard Station", "Exempted",
    ]
    transfer_categories = [
        "Mandatory JRTP Rotation", "Hard Station Replacement", "Administrative Rebalancing",
        "Promotion Posting", "Compassionate Grounds",
    ]
    approval_statuses = [
        "Recommended by DPC", "Pending Director (HR) Approval", "Approved by CMD",
        "Deferred", "Rejected",
    ]

    ws2.append(["Grade", "Discipline", "Tenure Status", "Hard Station Status", "Transfer Category", "Approval Status"])
    max_len = max(len(grades), len(disciplines), len(tenure_statuses), len(hard_station_statuses), len(transfer_categories), len(approval_statuses))
    for r in range(max_len):
        row_vals = [
            grades[r] if r < len(grades) else None,
            disciplines[r] if r < len(disciplines) else None,
            tenure_statuses[r] if r < len(tenure_statuses) else None,
            hard_station_statuses[r] if r < len(hard_station_statuses) else None,
            transfer_categories[r] if r < len(transfer_categories) else None,
            approval_statuses[r] if r < len(approval_statuses) else None,
        ]
        ws2.append(row_vals)

    # Style sheet 2 header
    hdr_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    hdr_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for cell in ws2[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font

    # Set up Sheet 1: Header banner
    ws1.merge_cells("A1:P1")
    title_cell = ws1["A1"]
    title_cell.value = "OIL AND NATURAL GAS CORPORATION LIMITED - ANNUAL GENERAL TRANSFER (AGT) PROPOSAL FY27"
    title_cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:P2")
    sub_cell = ws1["A2"]
    sub_cell.value = "Strictly Confidential | Prepared as per Job Rotation & Transfer Policy (JRTP) Rules | For DPC & Competent Authority Review"
    sub_cell.font = Font(name="Arial", size=9, italic=True, color="333333")
    sub_cell.fill = PatternFill(start_color="EAEFF5", end_color="EAEFF5", fill_type="solid")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # Headers for AGT table
    agt_headers = [
        "Proposal ID", "Employee ID", "Employee Name", "Grade", "Discipline",
        "Present Work Centre", "Present Station Tenure (Yrs)", "JRTP Tenure Status",
        "Proposed Work Centre", "Proposed Sector / Cluster", "Hard Station Requirement",
        "Transfer Category", "Replacement Proposed", "Estimated Transfer Cost (INR)",
        "Approval Status", "Administrative Justification & Remarks",
    ]
    ws1.append([])  # row 3 empty
    ws1.append(agt_headers)  # row 4
    ws1.row_dimensions[4].height = 24

    tbl_hdr_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    tbl_hdr_font = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    for col_idx, cell in enumerate(ws1[4], 1):
        cell.fill = tbl_hdr_fill
        cell.font = tbl_hdr_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Sample template row with placeholders
    sample_row = [
        "{{proposal_id}}", "{{emp_id}}", "{{employee_name}}", "{{grade_id}}", "{{discipline_id}}",
        "{{present_wc_id}}", "{{station_tenure_years}}", "{{jrtp_tenure_status}}",
        "{{proposed_wc_id}}", "{{proposed_cluster}}", "{{hard_station_status}}",
        "{{transfer_category}}", "{{replacement_emp_id}}", "{{transfer_cost_inr}}",
        "{{approval_status}}", "{{remarks}}",
    ]
    ws1.append(sample_row)  # row 5
    for cell in ws1[5]:
        cell.font = Font(name="Arial", size=9)
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="left", vertical="center")

    # Add Data Validations referencing Sheet 2
    dv_grade = DataValidation(type="list", formula1="Validation_Lists!$A$2:$A$11", allow_blank=True)
    dv_disc = DataValidation(type="list", formula1="Validation_Lists!$B$2:$B$18", allow_blank=True)
    dv_tenure = DataValidation(type="list", formula1="Validation_Lists!$C$2:$C$6", allow_blank=True)
    dv_hard = DataValidation(type="list", formula1="Validation_Lists!$D$2:$D$5", allow_blank=True)
    dv_cat = DataValidation(type="list", formula1="Validation_Lists!$E$2:$E$6", allow_blank=True)
    dv_appr = DataValidation(type="list", formula1="Validation_Lists!$F$2:$F$6", allow_blank=True)

    ws1.add_data_validation(dv_grade)
    ws1.add_data_validation(dv_disc)
    ws1.add_data_validation(dv_tenure)
    ws1.add_data_validation(dv_hard)
    ws1.add_data_validation(dv_cat)
    ws1.add_data_validation(dv_appr)

    dv_grade.add("D5:D500")
    dv_disc.add("E5:E500")
    dv_tenure.add("H5:H500")
    dv_hard.add("K5:K500")
    dv_cat.add("L5:L500")
    dv_appr.add("O5:O500")

    # Auto-fit column widths dynamically based on content length
    for col in ws1.columns:
        max_l = max(len(str(c.value or '')) for c in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = min(max(max_l + 2, 10), 40)

    wb.save(out_path)


def generate_parliament_reply_template(out_path: Path):
    """Generate the Parliament Question Reply template (DOCX with placeholders)."""
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hdr = p_hdr.add_run("GOVERNMENT OF INDIA\nMINISTRY OF PETROLEUM AND NATURAL GAS\n")
    r_hdr.font.name = "Arial"
    r_hdr.font.size = Pt(12)
    r_hdr.font.bold = True

    r_house = p_hdr.add_run("{{parliament_house}}\n")
    r_house.font.name = "Arial"
    r_house.font.size = Pt(13)
    r_house.font.bold = True
    r_house.font.color.rgb = RGBColor(0, 51, 102)

    r_sub = p_hdr.add_run("{{parliament_session}} SESSION | {{question_type}} QUESTION NO. {{question_number}}\nTO BE ANSWERED ON {{answer_date}}\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True

    # Subject box
    p_subj = doc.add_paragraph()
    p_subj.paragraph_format.space_before = Pt(8)
    p_subj.paragraph_format.space_after = Pt(8)
    r_s_lbl = p_subj.add_run("SUBJECT: ")
    r_s_lbl.font.bold = True
    p_subj.add_run("{{subject}}")

    # Member names
    p_mp = doc.add_paragraph()
    r_mp_lbl = p_mp.add_run("ASKED BY: ")
    r_mp_lbl.font.bold = True
    p_mp.add_run("{{mp_names}}")

    # Questions section
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    h_q = doc.add_heading("QUESTION", level=2)
    h_q.style.font.name = "Arial"
    h_q.style.font.color.rgb = RGBColor(0, 51, 102)

    p_q = doc.add_paragraph(
        "Will the Minister of PETROLEUM AND NATURAL GAS be pleased to state:\n\n"
        "(a) {{question_part_a}}\n"
        "(b) {{question_part_b}}\n"
        "(c) {{question_part_c}}\n"
        "(d) {{question_part_d}}"
    )
    p_q.paragraph_format.line_spacing = 1.15

    # Answer section
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    h_ans = doc.add_heading("ANSWER", level=2)
    h_ans.style.font.name = "Arial"
    h_ans.style.font.color.rgb = RGBColor(0, 51, 102)

    p_min = doc.add_paragraph()
    r_min = p_min.add_run("MINISTER OF STATE IN THE MINISTRY OF PETROLEUM AND NATURAL GAS\n({{minister_name}})\n")
    r_min.font.bold = True
    r_min.font.size = Pt(10)

    p_stmt = doc.add_paragraph(
        "(a) to (d): A statement is laid on the Table of the House.\n\n"
        "STATEMENT REFERRED TO IN REPLY TO PARTS (a) TO (d) OF {{parliament_house}} {{question_type}} QUESTION NO. {{question_number}} "
        "FOR {{answer_date}} REGARDING '{{subject}}'."
    )
    p_stmt.paragraph_format.line_spacing = 1.15

    # Text for parts a & b
    p_ab = doc.add_paragraph("{{answer_text_a_b}}")
    p_ab.paragraph_format.line_spacing = 1.15

    # Annexure Vacancy Table
    doc.add_paragraph().paragraph_format.space_before = Pt(4)
    p_tbl_lbl = doc.add_paragraph("TABLE: Sanctioned Strength, In-Position & Category-Wise Vacancies in ONGC:")
    p_tbl_lbl.runs[0].font.bold = True
    p_tbl_lbl.runs[0].font.size = Pt(9.5)

    vac_tbl = doc.add_table(rows=2, cols=9)
    vac_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    vac_headers = ["Work Centre / Asset", "Sanctioned Posts", "In-Position", "Total Vacancy", "SC", "ST", "OBC", "EWS", "General (UR)"]
    for c_idx, h in enumerate(vac_headers):
        c = vac_tbl.rows[0].cells[c_idx]
        c.paragraphs[0].text = h
        c.paragraphs[0].paragraph_format.space_before = Pt(3)
        c.paragraphs[0].paragraph_format.space_after = Pt(3)
        set_cell_background(c, "003366")
        for r in c.paragraphs[0].runs:
            r.font.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(c, 40, 40, 60, 60)

    sample_vac_row = [
        "{{wc_name}}", "{{sanctioned_total}}", "{{in_position_total}}", "{{vacancy_total}}",
        "{{sc_vac}}", "{{st_vac}}", "{{obc_vac}}", "{{ews_vac}}", "{{gen_vac}}",
    ]
    for c_idx, val in enumerate(sample_vac_row):
        c = vac_tbl.rows[1].cells[c_idx]
        c.paragraphs[0].text = val
        c.paragraphs[0].paragraph_format.space_before = Pt(3)
        c.paragraphs[0].paragraph_format.space_after = Pt(3)
        set_cell_background(c, "FFFFFF")
        for r in c.paragraphs[0].runs:
            r.font.size = Pt(8)
        set_cell_margins(c, 40, 40, 60, 60)

    # Text for parts c & d
    doc.add_paragraph().paragraph_format.space_before = Pt(6)
    p_cd = doc.add_paragraph("{{answer_text_c_d}}")
    p_cd.paragraph_format.line_spacing = 1.15

    # Verification / Vetting footer
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_foot = doc.add_paragraph()
    p_foot.add_run("Vetted by: {{vetted_by_official}}, Executive Director (HR), ONGC\n").font.bold = True
    p_foot.add_run("Approved by: {{approved_by_cmd}}, Chairman & Managing Director, ONGC\n").font.bold = True

    doc.save(out_path)


def generate_quarterly_asset_pnl_template(out_path: Path):
    """Generate the Quarterly Asset-Wise P&L template workbook (XLSX with placeholders)."""
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Quarterly_Asset_PnL"
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:G1")
    t_cell = ws["A1"]
    t_cell.value = "OIL AND NATURAL GAS CORPORATION LIMITED - QUARTERLY ASSET PERFORMANCE & OPERATING P&L"
    t_cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    t_cell.fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Metadata Row
    meta_pairs = [
        ("Asset Work Centre:", "{{asset_name}} ({{wc_id}})", "Sector / Basin:", "{{sector}} / {{basin_name}}"),
        ("Financial Quarter:", "FY {{fiscal_year}} - Quarter {{quarter}}", "Currency / Scale:", "INR (Crores)"),
    ]
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_pairs, 2):
        ws.cell(row=r_idx, column=1, value=k1).font = Font(name="Arial", size=9.5, bold=True)
        ws.cell(row=r_idx, column=2, value=v1).font = Font(name="Arial", size=9.5)
        ws.cell(row=r_idx, column=4, value=k2).font = Font(name="Arial", size=9.5, bold=True)
        ws.cell(row=r_idx, column=5, value=v2).font = Font(name="Arial", size=9.5)
        ws.row_dimensions[r_idx].height = 20

    # Table layout
    headers = ["Line Item / Head", "UOM", "Quarter Actual", "Previous Quarter", "Same Qtr Last Year", "Sanctioned BE", "Variance vs BE"]
    ws.append([])  # row 4 empty
    ws.append(headers)  # row 5
    ws.row_dimensions[5].height = 22

    th_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    th_font = Font(name="Arial", size=9.5, bold=True, color="FFFFFF")
    for cell in ws[5]:
        cell.fill = th_fill
        cell.font = th_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sec_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    sec_font = Font(name="Arial", size=9.5, bold=True, color="003366")
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    rows_def = [
        # Section A: Physical Production
        ("A. PHYSICAL HYDROCARBON PRODUCTION", "", "", "", "", "", ""),
        ("Crude Oil Production", "MT", "{{crude_oil_mt}}", "{{prev_qtr_oil_mt}}", "{{ly_qtr_oil_mt}}", "{{be_oil_mt}}", "{{var_oil_mt}}"),
        ("Natural Gas Production", "MMSCM", "{{natural_gas_mmscm}}", "{{prev_qtr_gas_mmscm}}", "{{ly_qtr_gas_mmscm}}", "{{be_gas_mmscm}}", "{{var_gas_mmscm}}"),
        ("Value Added Products / Condensate", "MT", "{{condensate_mt}}", "{{prev_qtr_cond_mt}}", "{{ly_qtr_cond_mt}}", "{{be_cond_mt}}", "{{var_cond_mt}}"),
        ("Total Production in Barrels of Oil Equivalent", "BOE", "{{total_boe}}", "{{prev_qtr_boe}}", "{{ly_qtr_boe}}", "{{be_boe}}", "{{var_boe}}"),
        ("Average Water Cut", "%", "{{water_cut_pct}}", "{{prev_qtr_water_cut}}", "{{ly_qtr_water_cut}}", "-", "-"),
        ("Active Producing Wells", "Count", "{{active_wells}}", "{{prev_qtr_wells}}", "{{ly_qtr_wells}}", "-", "-"),
        # Section B: Gross Revenue & Net Realization
        ("B. GROSS REVENUE & NET REALIZATION", "", "", "", "", "", ""),
        ("Crude Oil Sales Revenue", "₹ Cr", "{{crude_revenue_cr}}", "{{prev_crude_rev_cr}}", "{{ly_crude_rev_cr}}", "{{be_crude_rev_cr}}", "{{var_crude_rev_cr}}"),
        ("Natural Gas Sales Revenue", "₹ Cr", "{{gas_revenue_cr}}", "{{prev_gas_rev_cr}}", "{{ly_gas_rev_cr}}", "{{be_gas_rev_cr}}", "{{var_gas_rev_cr}}"),
        ("Condensate & VAP Sales", "₹ Cr", "{{vap_revenue_cr}}", "{{prev_vap_rev_cr}}", "{{ly_vap_rev_cr}}", "{{be_vap_rev_cr}}", "{{var_vap_rev_cr}}"),
        ("Total Gross Operating Revenue", "₹ Cr", "{{gross_revenue_cr}}", "{{prev_gross_rev_cr}}", "{{ly_gross_rev_cr}}", "{{be_gross_rev_cr}}", "{{var_gross_rev_cr}}"),
        ("Less: Royalty (Central / State)", "₹ Cr", "{{royalty_cr}}", "{{prev_royalty_cr}}", "{{ly_royalty_cr}}", "{{be_royalty_cr}}", "{{var_royalty_cr}}"),
        ("Less: Oil Industry Development (OID) Cess", "₹ Cr", "{{oid_cess_cr}}", "{{prev_cess_cr}}", "{{ly_cess_cr}}", "{{be_cess_cr}}", "{{var_cess_cr}}"),
        ("Net Operating Revenue", "₹ Cr", "{{net_revenue_cr}}", "{{prev_net_rev_cr}}", "{{ly_net_rev_cr}}", "{{be_net_rev_cr}}", "{{var_net_rev_cr}}"),
        # Section C: Operating Expenses (OPEX)
        ("C. OPERATING EXPENDITURE (OPEX)", "", "", "", "", "", ""),
        ("Manpower Costs (Salaries, Overtime, Medical, Benefits)", "₹ Cr", "{{manpower_opex_cr}}", "{{prev_manpower_cr}}", "{{ly_manpower_cr}}", "{{be_manpower_cr}}", "{{var_manpower_cr}}"),
        ("Offshore Platform / Surface Production Facilities", "₹ Cr", "{{facilities_opex_cr}}", "{{prev_facil_cr}}", "{{ly_facil_cr}}", "{{be_facil_cr}}", "{{var_facil_cr}}"),
        ("Workover & Well Servicing Operations", "₹ Cr", "{{workover_opex_cr}}", "{{prev_workover_cr}}", "{{ly_workover_cr}}", "{{be_workover_cr}}", "{{var_workover_cr}}"),
        ("Production Chemicals & Consumables", "₹ Cr", "{{chemicals_opex_cr}}", "{{prev_chem_cr}}", "{{ly_chem_cr}}", "{{be_chem_cr}}", "{{var_chem_cr}}"),
        ("Logistics, Offshore Supply Vessels & Helicopter Support", "₹ Cr", "{{logistics_opex_cr}}", "{{prev_logist_cr}}", "{{ly_logist_cr}}", "{{be_logist_cr}}", "{{var_logist_cr}}"),
        ("Power, Gas Fuel & Utilities", "₹ Cr", "{{utilities_opex_cr}}", "{{prev_util_cr}}", "{{ly_util_cr}}", "{{be_util_cr}}", "{{var_util_cr}}"),
        ("General Asset Administration & Overheads", "₹ Cr", "{{admin_opex_cr}}", "{{prev_admin_cr}}", "{{ly_admin_cr}}", "{{be_admin_cr}}", "{{var_admin_cr}}"),
        ("Total Operating Expenses (OPEX)", "₹ Cr", "{{total_opex_cr}}", "{{prev_opex_cr}}", "{{ly_opex_cr}}", "{{be_opex_cr}}", "{{var_opex_cr}}"),
        # Section D: Profitability & Unit Cost Ratios
        ("D. PROFITABILITY & UNIT FINANCIAL RATIOS", "", "", "", "", "", ""),
        ("Operating EBITDA", "₹ Cr", "{{ebitda_cr}}", "{{prev_ebitda_cr}}", "{{ly_ebitda_cr}}", "{{be_ebitda_cr}}", "{{var_ebitda_cr}}"),
        ("EBITDA Margin", "%", "{{ebitda_margin_pct}}", "{{prev_margin_pct}}", "{{ly_margin_pct}}", "{{be_margin_pct}}", "{{var_margin_pct}}"),
        ("Lifting Cost per BOE", "₹/BOE", "{{lifting_cost_inr_per_boe}}", "{{prev_lift_inr}}", "{{ly_lift_inr}}", "{{be_lift_inr}}", "{{var_lift_inr}}"),
        ("Lifting Cost per BOE", "$/BOE", "{{lifting_cost_usd_per_boe}}", "{{prev_lift_usd}}", "{{ly_lift_usd}}", "{{be_lift_usd}}", "{{var_lift_usd}}"),
        # Section E: Capex Performance
        ("E. CAPITAL EXPENDITURE (CAPEX)", "", "", "", "", "", ""),
        ("Sanctioned Capex Budget (BE)", "₹ Cr", "{{capex_be_cr}}", "{{prev_capex_be_cr}}", "{{ly_capex_be_cr}}", "{{be_capex_cr}}", "-"),
        ("Actual Capex Spend Incurred", "₹ Cr", "{{capex_actual_cr}}", "{{prev_capex_act_cr}}", "{{ly_capex_act_cr}}", "{{be_capex_act_cr}}", "{{var_capex_cr}}"),
        ("Capex Utilization / Burn Rate", "%", "{{capex_utilization_pct}}", "{{prev_burn_pct}}", "{{ly_burn_pct}}", "100.0%", "{{var_burn_pct}}"),
    ]

    for item in rows_def:
        ws.append(list(item))
        curr_row = ws.max_row
        ws.row_dimensions[curr_row].height = 19
        is_section_header = item[0].startswith(("A.", "B.", "C.", "D.", "E."))
        for col_idx, cell in enumerate(ws[curr_row], 1):
            cell.border = thin_border
            if is_section_header:
                cell.fill = sec_fill
                cell.font = sec_font
            else:
                cell.font = Font(name="Arial", size=9)
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_idx == 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")

    # Dynamic column widths
    for col in ws.columns:
        max_l = max(len(str(c.value or '')) for c in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_l + 2, 10), 55)

    # Sheet 2: Notes & Accounting Policies
    ws_notes = wb.create_sheet(title="Notes_and_Policies")
    ws_notes.append(["Reference Note", "Accounting & Operational Policy Guidelines"])
    ws_notes.append(["Lifting Cost", "Computed as Total OPEX divided by Gross Hydrocarbon Volume in BOE (using standard 1 MT crude = 7.33 boe; 1 MMSCM gas = 5,882 boe)."])
    ws_notes.append(["Statutory Levies", "Royalty rates as prescribed by MoPNG (20% onshore crude, 10% offshore crude deepwater exemptions applied as per PSC)."])
    ws_notes.append(["Placeholders", "All {{field_name}} tokens are resolved dynamically during quarterly review generation."])
    ws_notes.column_dimensions["A"].width = 20
    ws_notes.column_dimensions["B"].width = 80

    wb.save(out_path)


def main():
    TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)

    p_board = TEMPLATES_DIR / "template_board_note.docx"
    p_agt = TEMPLATES_DIR / "template_annual_general_transfer_proposal.xlsx"
    p_parl = TEMPLATES_DIR / "template_parliament_reply_format.docx"
    p_pnl = TEMPLATES_DIR / "template_quarterly_asset_pnl.xlsx"

    print("Generating Board Note Template...")
    generate_board_note_template(p_board)

    print("Generating AGT Proposal Template...")
    generate_agt_proposal_template(p_agt)

    print("Generating Parliament Reply Template...")
    generate_parliament_reply_template(p_parl)

    print("Generating Quarterly Asset P&L Template...")
    generate_quarterly_asset_pnl_template(p_pnl)

    print("All 4 templates generated successfully.")


if __name__ == "__main__":
    main()
