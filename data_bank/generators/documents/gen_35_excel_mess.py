"""gen_35_excel_mess.py: Post-processor introducing realistic ONGC Excel source messiness.

Transforms clean single-header Phase-2 Excel workbooks into authentic, messy layouts:
- Title blocks & division headers
- Merged multi-row category headers
- Unit rows ('₹ crore', 'in Nos.', etc.)
- Blank spacer columns
- Multi-sheet workbooks (per-asset or per-FY partitions)
- Interspersed subtotal / total rows
- Explanatory notes and confidentiality footers
- Hidden metadata / lookup sheets
- Formatted numeric strings with Indian comma formatting

Outputs:
- landing/excel/<workbook>.xlsx (messy source-native layout)
- landing/excel/_clean/<workbook>.xlsx (pristine originals preserved)
- landing/excel/_parsing_specs.yaml (declarative parsing rules for Phase 4)
"""

from __future__ import annotations

import os
import shutil
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import yaml

BASE_DIR = Path(__file__).resolve().parent.parent.parent
LANDING_EXCEL = BASE_DIR / "landing" / "excel"
CLEAN_EXCEL = LANDING_EXCEL / "_clean"


def format_indian_number(val: Any, decimals: int = 2) -> str:
    """Format numeric values with standard Indian numbering commas (Lakhs/Crores)."""
    if pd.isna(val) or val is None or val == "":
        return ""
    try:
        fval = float(val)
    except (ValueError, TypeError):
        return str(val)

    is_neg = fval < 0
    fval = abs(fval)
    s = f"{fval:.{decimals}f}"
    parts = s.split(".")
    integer_part = parts[0]
    dec_part = "." + parts[1] if decimals > 0 else ""

    if len(integer_part) <= 3:
        formatted = integer_part
    else:
        last3 = integer_part[-3:]
        rest = integer_part[:-3]
        groups = []
        while len(rest) > 2:
            groups.insert(0, rest[-2:])
            rest = rest[:-2]
        if rest:
            groups.insert(0, rest)
        formatted = ",".join(groups) + "," + last3

    return ("-" if is_neg else "") + formatted + dec_part


def ensure_clean_backups(target_files: List[str]) -> None:
    """Ensure clean originals exist under landing/excel/_clean/."""
    CLEAN_EXCEL.mkdir(parents=True, exist_ok=True)
    for fname in target_files:
        src = LANDING_EXCEL / fname
        dst = CLEAN_EXCEL / fname
        if not dst.exists() and src.exists():
            shutil.copy2(src, dst)


def apply_header_styles(ws: openpyxl.worksheet.worksheet.Worksheet, max_col: int) -> None:
    """Apply styling to title block, category headers, column headers, and unit row."""
    navy_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    dark_gray_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # Title row 1
    cell_t1 = ws.cell(1, 2)
    cell_t1.font = Font(name="Calibri", size=13, bold=True, color="1A365D")

    # Title row 2
    cell_t2 = ws.cell(2, 2)
    cell_t2.font = Font(name="Calibri", size=10, italic=True, color="4A5568")

    # Category header row 4
    for c in range(2, max_col + 1):
        cell = ws.cell(4, c)
        if cell.value:
            cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            cell.fill = navy_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # Column header row 5
    for c in range(2, max_col + 1):
        cell = ws.cell(5, c)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = dark_gray_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Unit row 6
    for c in range(2, max_col + 1):
        cell = ws.cell(6, c)
        cell.font = Font(name="Calibri", size=8, italic=True, color="2B6CB0")
        cell.fill = light_blue_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


def process_capex_be_re_actuals() -> None:
    """Rewrite capex_be_re_actuals.xlsx: BE/RE workbook per asset sheet + subtotals + Indian format."""
    clean_path = CLEAN_EXCEL / "capex_be_re_actuals.xlsx"
    df = pd.read_excel(clean_path)
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    # Create sheets per asset_wc_id
    asset_col = "asset_wc_id"
    groups = df.groupby(asset_col, sort=False)
    numeric_cols = ["budget_estimate_be_inr_cr", "revised_estimate_re_inr_cr", "actual_spend_inr_cr"]

    for asset_val, group in groups:
        ws = wb.create_sheet(title=str(asset_val))

        # Title block
        ws.cell(1, 2, value="OIL AND NATURAL GAS CORPORATION LIMITED - CAPITAL EXPENDITURE MONITORING")
        ws.cell(2, 2, value=f"Asset Expenditure Ledger: {asset_val} | BE / RE / Actuals Tracking")

        # Merged Category Headers (Row 4)
        ws.cell(4, 2, value="PROJECT IDENTIFICATION")
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=4)
        ws.cell(4, 5, value="PERIOD & CLASSIFICATION")
        ws.merge_cells(start_row=4, start_column=5, end_row=4, end_column=8)
        ws.cell(4, 9, value="CAPEX EXPENDITURE (₹ CRORE)")
        ws.merge_cells(start_row=4, start_column=9, end_row=4, end_column=11)

        # Column Headers (Row 5) - starting at Col B (Col A is blank spacer)
        cols = list(df.columns)
        for col_idx, col_id in enumerate(cols, start=2):
            ws.cell(5, col_idx, value=col_id)

        # Unit Row (Row 6)
        unit_map = {
            "budget_estimate_be_inr_cr": "(₹ crore)",
            "revised_estimate_re_inr_cr": "(₹ crore)",
            "actual_spend_inr_cr": "(₹ crore)",
        }
        for col_idx, col_id in enumerate(cols, start=2):
            ws.cell(6, col_idx, value=unit_map.get(col_id, ""))

        # Data Rows (Row 7+)
        current_row = 7
        for _, row in group.iterrows():
            for col_idx, col_id in enumerate(cols, start=2):
                val = row[col_id]
                if col_id in numeric_cols and pd.notna(val):
                    # Store some numbers with Indian comma format as text
                    if current_row % 2 == 0:
                        cell_val = format_indian_number(val, 2)
                    else:
                        cell_val = float(val)
                elif pd.isna(val):
                    cell_val = None
                else:
                    cell_val = val
                ws.cell(current_row, col_idx, value=cell_val)
            current_row += 1

        # Subtotal row
        ws.cell(current_row, 2, value="Subtotal - Asset Expenditure Total")
        for col_idx, col_id in enumerate(cols, start=2):
            if col_id in numeric_cols:
                sum_val = group[col_id].sum()
                ws.cell(current_row, col_idx, value=format_indian_number(sum_val, 2))
        ws.row_dimensions[current_row].height = 20
        for c in range(2, len(cols) + 2):
            cell = ws.cell(current_row, c)
            cell.font = Font(name="Calibri", size=9, bold=True, color="1A202C")
            cell.fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        current_row += 1

        # Notes / Footer
        ws.cell(current_row, 2, value="Note: Compiled from SAP PS/CO records. Figures subject to statutory reconciliation.")
        ws.cell(current_row, 2).font = Font(name="Calibri", size=8, italic=True, color="718096")

        # Styling
        apply_header_styles(ws, len(cols) + 1)
        ws.column_dimensions["A"].width = 3

    # Remove the initial default empty sheet
    wb.remove(default_sheet)

    # Hidden Sheet
    hidden_ws = wb.create_sheet(title="_Ref_Capex_Codes")
    hidden_ws.sheet_state = "hidden"
    hidden_ws.cell(1, 1, value="config_key")
    hidden_ws.cell(1, 2, value="config_value")
    hidden_ws.cell(2, 1, value="version")
    hidden_ws.cell(2, 2, value="FY17-FY26-v2.1")
    hidden_ws.cell(3, 1, value="source_system")
    hidden_ws.cell(3, 2, value="SAP_CO_PA_PS")

    out_path = LANDING_EXCEL / "capex_be_re_actuals.xlsx"
    wb.save(out_path)
    print(f"Saved messy {out_path.name}: {len(wb.sheetnames)} sheets")


def process_manpower_plan() -> None:
    """Rewrite manpower_plan.xlsx: title, merged headers, unit row, spacer, subtotals, hidden sheet."""
    clean_path = CLEAN_EXCEL / "manpower_plan.xlsx"
    df = pd.read_excel(clean_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "manpower_plan"

    # Title block
    ws.cell(1, 2, value="OIL AND NATURAL GAS CORPORATION LIMITED - CORPORATE HR")
    ws.cell(2, 2, value="ANNUAL SANCTIONED POSTS & MANPOWER ROLLOUT PLAN - FY 2025-26")

    # Merged Category Headers (Row 4)
    ws.cell(4, 2, value="CADRE IDENTIFIERS")
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=5)
    ws.cell(4, 6, value="CADRE & DEMOGRAPHICS")
    ws.merge_cells(start_row=4, start_column=6, end_row=4, end_column=12)
    ws.cell(4, 13, value="STRENGTH & VACANCIES")
    ws.merge_cells(start_row=4, start_column=13, end_row=4, end_column=17)
    ws.cell(4, 18, value="SPECIAL RECRUITMENT DRIVES")
    ws.merge_cells(start_row=4, start_column=18, end_row=4, end_column=20)

    # Column Headers (Row 5)
    cols = list(df.columns)
    for col_idx, col_id in enumerate(cols, start=2):
        ws.cell(5, col_idx, value=col_id)

    # Unit Row (Row 6)
    unit_map = {
        "sanctioned_strength": "(in Nos.)",
        "in_position_strength": "(in Nos.)",
        "vacancy_count": "(in Nos.)",
        "vacancy_rate_pct": "(Percentage)",
        "backlog_vacancies": "(in Nos.)",
        "special_drive_flag": "(Y/N)",
        "special_drive_target_vacancies": "(in Nos.)",
    }
    for col_idx, col_id in enumerate(cols, start=2):
        ws.cell(6, col_idx, value=unit_map.get(col_id, ""))

    # Data Rows (Row 7+)
    int_cols = ["sanctioned_strength", "in_position_strength", "vacancy_count", "backlog_vacancies"]
    current_row = 7
    for _, row in df.iterrows():
        for col_idx, col_id in enumerate(cols, start=2):
            val = row[col_id]
            if col_id in int_cols and pd.notna(val) and int(val) > 999 and current_row % 3 == 0:
                cell_val = format_indian_number(val, 0)
            elif pd.isna(val):
                cell_val = None
            else:
                cell_val = val
            ws.cell(current_row, col_idx, value=cell_val)
        current_row += 1

    # Subtotal row
    ws.cell(current_row, 2, value="Total - Corporate Strength Summary")
    for col_idx, col_id in enumerate(cols, start=2):
        if col_id in int_cols:
            ws.cell(current_row, col_idx, value=format_indian_number(df[col_id].sum(), 0))
    ws.row_dimensions[current_row].height = 20
    for c in range(2, len(cols) + 2):
        cell = ws.cell(current_row, c)
        cell.font = Font(name="Calibri", size=9, bold=True, color="1A202C")
        cell.fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    current_row += 1

    # Notes / Footer
    ws.cell(current_row, 2, value="Note: Sanctioned cadre limits as approved by 342nd ONGC Board Meeting. Strictly for internal HR planning.")
    ws.cell(current_row, 2).font = Font(name="Calibri", size=8, italic=True, color="718096")

    # Styling
    apply_header_styles(ws, len(cols) + 1)
    ws.column_dimensions["A"].width = 3

    # Hidden Sheet
    hidden_ws = wb.create_sheet(title="_Cadre_Control_Master")
    hidden_ws.sheet_state = "hidden"
    hidden_ws.cell(1, 1, value="key")
    hidden_ws.cell(1, 2, value="value")
    hidden_ws.cell(2, 1, value="plan_year")
    hidden_ws.cell(2, 2, value="2025-2026")
    hidden_ws.cell(3, 1, value="board_approval_date")
    hidden_ws.cell(3, 2, value="2025-03-24")

    out_path = LANDING_EXCEL / "manpower_plan.xlsx"
    wb.save(out_path)
    print(f"Saved messy {out_path.name}")


def process_dpe_mou_scorecard() -> None:
    """Rewrite dpe_mou_scorecard.xlsx: per-FY sheets, multi-row headers, unit row, subtotals, hidden sheet."""
    clean_path = CLEAN_EXCEL / "dpe_mou_scorecard.xlsx"
    df = pd.read_excel(clean_path)
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    fy_col = "fiscal_year"
    groups = df.groupby(fy_col, sort=False)
    cols = list(df.columns)
    numeric_cols = [
        "weightage",
        "excellent_target",
        "very_good_target",
        "good_target",
        "fair_target",
        "poor_target",
        "actual_achievement",
        "weighted_score",
        "composite_score",
        "projected_composite_score",
    ]

    for fy_val, group in groups:
        ws = wb.create_sheet(title=str(fy_val))

        # Title block
        ws.cell(1, 2, value="DEPARTMENT OF PUBLIC ENTERPRISES (DPE) - ANNUAL MOU PERFORMANCE SCORECARD")
        ws.cell(2, 2, value=f"ONGC Target Evaluation & Composite Scorecard - {fy_val}")

        # Category Headers (Row 4)
        ws.cell(4, 2, value="MOU PARAMETER DETAILS")
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=7)
        ws.cell(4, 8, value="DPE TARGET CRITERIA")
        ws.merge_cells(start_row=4, start_column=8, end_row=4, end_column=13)
        ws.cell(4, 14, value="PERFORMANCE & SCORING")
        ws.merge_cells(start_row=4, start_column=14, end_row=4, end_column=18)
        ws.cell(4, 19, value="REMARKS")

        # Column Headers (Row 5)
        for col_idx, col_id in enumerate(cols, start=2):
            ws.cell(5, col_idx, value=col_id)

        # Unit Row (Row 6)
        unit_map = {
            "weightage": "(Weight %)",
            "excellent_target": "(Target)",
            "very_good_target": "(Target)",
            "good_target": "(Target)",
            "fair_target": "(Target)",
            "poor_target": "(Target)",
            "actual_achievement": "(Actual)",
            "weighted_score": "(Score)",
            "composite_score": "(Composite)",
            "projected_composite_score": "(Projected)",
        }
        for col_idx, col_id in enumerate(cols, start=2):
            ws.cell(6, col_idx, value=unit_map.get(col_id, ""))

        # Data rows (Row 7+)
        target_eval_cols = ["excellent_target", "actual_achievement"]
        current_row = 7
        for _, row in group.iterrows():
            for col_idx, col_id in enumerate(cols, start=2):
                val = row[col_id]
                if col_id in target_eval_cols and pd.notna(val) and current_row % 2 == 0:
                    cell_val = format_indian_number(val, 2)
                elif pd.isna(val):
                    cell_val = None
                else:
                    cell_val = val
                ws.cell(current_row, col_idx, value=cell_val)
            current_row += 1

        # Subtotal row
        ws.cell(current_row, 2, value="Total / Composite Summary")
        for col_idx, col_id in enumerate(cols, start=2):
            if col_id == "weightage":
                ws.cell(current_row, col_idx, value=float(group["weightage"].sum()))
            elif col_id == "weighted_score":
                ws.cell(current_row, col_idx, value=round(float(group["weighted_score"].sum()), 4))
        ws.row_dimensions[current_row].height = 20
        for c in range(2, len(cols) + 2):
            cell = ws.cell(current_row, c)
            cell.font = Font(name="Calibri", size=9, bold=True, color="1A202C")
            cell.fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        current_row += 1

        # Footer note
        ws.cell(current_row, 2, value="Status: Evaluated per DPE Guidelines OM No. DPE/17(1)/2019-Fin. Approved by Inter-Ministerial Committee.")
        ws.cell(current_row, 2).font = Font(name="Calibri", size=8, italic=True, color="718096")

        apply_header_styles(ws, len(cols) + 1)
        ws.column_dimensions["A"].width = 3

    wb.remove(default_sheet)

    # Hidden Sheet
    hidden_ws = wb.create_sheet(title="_MoU_Grading_Rules")
    hidden_ws.sheet_state = "hidden"
    hidden_ws.cell(1, 1, value="rating")
    hidden_ws.cell(1, 2, value="composite_score_range")
    hidden_ws.cell(2, 1, value="Excellent")
    hidden_ws.cell(2, 2, value="1.00 - 1.50")
    hidden_ws.cell(3, 1, value="Very Good")
    hidden_ws.cell(3, 2, value="1.51 - 2.50")

    out_path = LANDING_EXCEL / "dpe_mou_scorecard.xlsx"
    wb.save(out_path)
    print(f"Saved messy {out_path.name}: {len(wb.sheetnames)} sheets")


def process_csr_projects_tracker() -> None:
    """Rewrite csr_projects_tracker.xlsx: per-FY sheets, multi-row headers, unit row, subtotals, hidden sheet."""
    clean_path = CLEAN_EXCEL / "csr_projects_tracker.xlsx"
    df = pd.read_excel(clean_path)
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    fy_col = "fiscal_year"
    groups = df.groupby(fy_col, sort=False)
    cols = list(df.columns)
    numeric_cols = ["sanctioned_budget_inr_cr", "actual_disbursed_inr_cr", "beneficiaries_count"]

    for fy_val, group in groups:
        ws = wb.create_sheet(title=str(fy_val))

        # Title block
        ws.cell(1, 2, value="ONGC FOUNDATION - STATUTORY CSR ALLOCATIONS & DISBURSEMENTS")
        ws.cell(2, 2, value=f"Annual CSR Implementation Progress Schedule - {fy_val}")

        # Category Headers (Row 4)
        ws.cell(4, 2, value="PROJECT & LOCATION")
        ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=6)
        ws.cell(4, 7, value="THEMATIC & JURISDICTION")
        ws.merge_cells(start_row=4, start_column=7, end_row=4, end_column=10)
        ws.cell(4, 11, value="FINANCIAL PROGRESS (₹ CRORE)")
        ws.merge_cells(start_row=4, start_column=11, end_row=4, end_column=12)
        ws.cell(4, 13, value="IMPLEMENTATION & BENEFICIARIES")
        ws.merge_cells(start_row=4, start_column=13, end_row=4, end_column=15)

        # Column Headers (Row 5)
        for col_idx, col_id in enumerate(cols, start=2):
            ws.cell(5, col_idx, value=col_id)

        # Unit Row (Row 6)
        unit_map = {
            "sanctioned_budget_inr_cr": "(₹ crore)",
            "actual_disbursed_inr_cr": "(₹ crore)",
            "beneficiaries_count": "(in Nos.)",
            "statutory_2pct_mandate_met_flag": "(Y/N)",
        }
        for col_idx, col_id in enumerate(cols, start=2):
            ws.cell(6, col_idx, value=unit_map.get(col_id, ""))

        # Data rows (Row 7+)
        spend_cols = ["sanctioned_budget_inr_cr", "actual_disbursed_inr_cr"]
        current_row = 7
        for _, row in group.iterrows():
            for col_idx, col_id in enumerate(cols, start=2):
                val = row[col_id]
                if col_id in spend_cols and pd.notna(val) and current_row % 2 == 0:
                    cell_val = format_indian_number(val, 2)
                elif col_id == "beneficiaries_count" and pd.notna(val) and int(val) > 999:
                    cell_val = format_indian_number(val, 0)
                elif pd.isna(val):
                    cell_val = None
                else:
                    cell_val = val
                ws.cell(current_row, col_idx, value=cell_val)
            current_row += 1

        # Subtotal row
        ws.cell(current_row, 2, value="Subtotal - Annual CSR Spend Total")
        for col_idx, col_id in enumerate(cols, start=2):
            if col_id in spend_cols:
                ws.cell(current_row, col_idx, value=format_indian_number(group[col_id].sum(), 2))
            elif col_id == "beneficiaries_count":
                ws.cell(current_row, col_idx, value=format_indian_number(group[col_id].sum(), 0))
        ws.row_dimensions[current_row].height = 20
        for c in range(2, len(cols) + 2):
            cell = ws.cell(current_row, c)
            cell.font = Font(name="Calibri", size=9, bold=True, color="1A202C")
            cell.fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        current_row += 1

        # Footer note
        ws.cell(current_row, 2, value="Mandate Note: 2% of average net profits of preceding 3 financial years pursuant to Section 135 of Companies Act, 2013.")
        ws.cell(current_row, 2).font = Font(name="Calibri", size=8, italic=True, color="718096")

        apply_header_styles(ws, len(cols) + 1)
        ws.column_dimensions["A"].width = 3

    wb.remove(default_sheet)

    # Hidden Sheet
    hidden_ws = wb.create_sheet(title="_Schedule_VII_Themes")
    hidden_ws.sheet_state = "hidden"
    hidden_ws.cell(1, 1, value="theme_code")
    hidden_ws.cell(1, 2, value="statutory_description")
    hidden_ws.cell(2, 1, value="THEME-EDU")
    hidden_ws.cell(2, 2, value="Promoting education and vocational training")
    hidden_ws.cell(3, 1, value="THEME-HLT")
    hidden_ws.cell(3, 2, value="Promoting preventive healthcare and sanitation")

    out_path = LANDING_EXCEL / "csr_projects_tracker.xlsx"
    wb.save(out_path)
    print(f"Saved messy {out_path.name}: {len(wb.sheetnames)} sheets")


def process_succession_planning() -> None:
    """Rewrite succession_planning.xlsx: title block, category headers, unit row, spacer, subtotal, hidden sheet."""
    clean_path = CLEAN_EXCEL / "succession_planning.xlsx"
    df = pd.read_excel(clean_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "succession_plan"

    # Title block
    ws.cell(1, 2, value="CONFIDENTIAL - BOARD OF DIRECTORS & KEY MANAGEMENT SUCCESSION")
    ws.cell(2, 2, value="CRITICAL LEADERSHIP PIPELINE & TALENT BENCH AUDIT - FY26")

    # Category Headers (Row 4)
    ws.cell(4, 2, value="KEY POSITION IDENTIFICATION")
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=6)
    ws.cell(4, 7, value="INCUMBENT PROFILE")
    ws.merge_cells(start_row=4, start_column=7, end_row=4, end_column=9)
    ws.cell(4, 10, value="PRIMARY SUCCESSOR (BENCH 1)")
    ws.merge_cells(start_row=4, start_column=10, end_row=4, end_column=12)
    ws.cell(4, 13, value="SECONDARY SUCCESSOR (BENCH 2)")
    ws.merge_cells(start_row=4, start_column=13, end_row=4, end_column=15)
    ws.cell(4, 16, value="COMPETENCY GAP")

    # Column Headers (Row 5)
    cols = list(df.columns)
    for col_idx, col_id in enumerate(cols, start=2):
        ws.cell(5, col_idx, value=col_id)

    # Unit Row (Row 6)
    unit_map = {
        "position_id": "(ID)",
        "role_title": "(Designation)",
        "criticality_tier": "(Tier)",
        "incumbent_retire_date": "(YYYY-MM-DD)",
        "successor_1_has_ne_hard_station": "(Y/N)",
        "successor_2_has_ne_hard_station": "(Y/N)",
    }
    for col_idx, col_id in enumerate(cols, start=2):
        ws.cell(6, col_idx, value=unit_map.get(col_id, ""))

    # Data Rows (Row 7+)
    current_row = 7
    for _, row in df.iterrows():
        for col_idx, col_id in enumerate(cols, start=2):
            val = row[col_id]
            cell_val = None if pd.isna(val) else val
            ws.cell(current_row, col_idx, value=cell_val)
        current_row += 1

    # Subtotal row
    ws.cell(current_row, 2, value=f"Total Key Positions Assessed: {len(df)}")
    for c in range(2, len(cols) + 2):
        cell = ws.cell(current_row, c)
        cell.font = Font(name="Calibri", size=9, bold=True, color="1A202C")
        cell.fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    current_row += 1

    # Footer note
    ws.cell(current_row, 2, value="Strictly Confidential. Subject to DPC and Search-cum-Selection Committee (PESB) regulations.")
    ws.cell(current_row, 2).font = Font(name="Calibri", size=8, italic=True, color="718096")

    apply_header_styles(ws, len(cols) + 1)
    ws.column_dimensions["A"].width = 3

    # Hidden Sheet
    hidden_ws = wb.create_sheet(title="_Talent_Bench_Scale")
    hidden_ws.sheet_state = "hidden"
    hidden_ws.cell(1, 1, value="readiness_tier")
    hidden_ws.cell(1, 2, value="expected_tenure_horizon")
    hidden_ws.cell(2, 1, value="Ready_Now")
    hidden_ws.cell(2, 2, value="< 6 months")
    hidden_ws.cell(3, 1, value="Ready_1_2_Years")
    hidden_ws.cell(3, 2, value="12 - 24 months")

    out_path = LANDING_EXCEL / "succession_planning.xlsx"
    wb.save(out_path)
    print(f"Saved messy {out_path.name}")


def process_legal_cases_tracker() -> None:
    """Rewrite legal_cases_tracker.xlsx: title, headers, unit row, spacer, subtotals, hidden sheet."""
    clean_path = CLEAN_EXCEL / "legal_cases_tracker.xlsx"
    df = pd.read_excel(clean_path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "legal_cases"

    # Title block
    ws.cell(1, 2, value="OIL AND NATURAL GAS CORPORATION LIMITED - CORPORATE LEGAL DIVISION")
    ws.cell(2, 2, value="COURT CASES & DISCIPLINARY LITIGATION TRACKER - STATUS REGISTER")

    # Category Headers (Row 4)
    ws.cell(4, 2, value="CASE PARTICULARS")
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=6)
    ws.cell(4, 7, value="FORUM & FILING")
    ws.merge_cells(start_row=4, start_column=7, end_row=4, end_column=9)
    ws.cell(4, 10, value="STATUS & STAY")
    ws.merge_cells(start_row=4, start_column=10, end_row=4, end_column=11)
    ws.cell(4, 12, value="AGING ANALYSIS")
    ws.merge_cells(start_row=4, start_column=12, end_row=4, end_column=14)
    ws.cell(4, 15, value="HEARING & BRIEF")
    ws.merge_cells(start_row=4, start_column=15, end_row=4, end_column=16)

    # Column Headers (Row 5)
    cols = list(df.columns)
    for col_idx, col_id in enumerate(cols, start=2):
        ws.cell(5, col_idx, value=col_id)

    # Unit Row (Row 6)
    unit_map = {
        "filing_date": "(YYYY-MM-DD)",
        "interim_stay_granted": "(Y/N)",
        "age_days": "(Days)",
        "age_years": "(Years)",
        "is_aged_over_3_years": "(Y/N)",
        "next_hearing_date": "(YYYY-MM-DD)",
    }
    for col_idx, col_id in enumerate(cols, start=2):
        ws.cell(6, col_idx, value=unit_map.get(col_id, ""))

    # Data Rows (Row 7+)
    current_row = 7
    for _, row in df.iterrows():
        for col_idx, col_id in enumerate(cols, start=2):
            val = row[col_id]
            if col_id == "age_days" and pd.notna(val) and int(val) > 999 and current_row % 2 == 0:
                cell_val = format_indian_number(val, 0)
            elif col_id == "age_years" and pd.notna(val) and current_row % 3 == 0:
                cell_val = format_indian_number(val, 2)
            elif pd.isna(val):
                cell_val = None
            else:
                cell_val = val
            ws.cell(current_row, col_idx, value=cell_val)
        current_row += 1

    # Subtotal row
    ws.cell(current_row, 2, value=f"Total - Active Litigation Cases: {len(df)}")
    for c in range(2, len(cols) + 2):
        cell = ws.cell(current_row, c)
        cell.font = Font(name="Calibri", size=9, bold=True, color="1A202C")
        cell.fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    current_row += 1

    # Footer note
    ws.cell(current_row, 2, value="Privileged & Confidential. For Legal Department and Senior Management use only.")
    ws.cell(current_row, 2).font = Font(name="Calibri", size=8, italic=True, color="718096")

    apply_header_styles(ws, len(cols) + 1)
    ws.column_dimensions["A"].width = 3

    # Hidden Sheet
    hidden_ws = wb.create_sheet(title="_Court_Jurisdiction_Codes")
    hidden_ws.sheet_state = "hidden"
    hidden_ws.cell(1, 1, value="code")
    hidden_ws.cell(1, 2, value="description")
    hidden_ws.cell(2, 1, value="CAT")
    hidden_ws.cell(2, 2, value="Central Administrative Tribunal")
    hidden_ws.cell(3, 1, value="SC")
    hidden_ws.cell(3, 2, value="Supreme Court of India")

    out_path = LANDING_EXCEL / "legal_cases_tracker.xlsx"
    wb.save(out_path)
    print(f"Saved messy {out_path.name}")


def generate_parsing_specs() -> None:
    """Generate declarative _parsing_specs.yaml for Phase 4 Excel loading."""
    specs = {
        "version": "1.0",
        "description": "Parsing specifications for realistic messy ONGC Excel sources (Step 3.5)",
        "workbooks": {
            "capex_be_re_actuals.xlsx": {
                "target_table": "capex_budget",
                "sheet_strategy": {
                    "mode": "all_visible",
                    "exclude_sheets": ["_Ref_Capex_Codes"],
                },
                "header_row": 5,
                "unit_row": 6,
                "data_start_row": 7,
                "drop_columns": ["Unnamed: 0"],
                "subtotal_filter": {
                    "column": "budget_id",
                    "patterns": ["^Subtotal", "^Total", "^Summary"],
                },
                "footer_filter": {
                    "column": "budget_id",
                    "patterns": ["^Note:", "^Confidential", "^Status:"],
                },
                "numeric_columns": {
                    "budget_estimate_be_inr_cr": "float",
                    "revised_estimate_re_inr_cr": "float",
                    "actual_spend_inr_cr": "float",
                },
                "boolean_columns": [],
                "integer_columns": [],
                "string_columns": [
                    "budget_id",
                    "prj_id",
                    "project_name",
                    "asset_wc_id",
                    "month",
                    "fiscal_year",
                    "capex_head",
                ],
                "primary_key": ["budget_id"],
            },
            "manpower_plan.xlsx": {
                "target_table": "sanctioned_posts",
                "sheet_strategy": {
                    "mode": "named",
                    "sheet_name": "manpower_plan",
                    "exclude_sheets": ["_Cadre_Control_Master"],
                },
                "header_row": 5,
                "unit_row": 6,
                "data_start_row": 7,
                "drop_columns": ["Unnamed: 0"],
                "subtotal_filter": {
                    "column": "sanction_id",
                    "patterns": ["^Subtotal", "^Total", "^Summary"],
                },
                "footer_filter": {
                    "column": "sanction_id",
                    "patterns": ["^Note:", "^Confidential"],
                },
                "numeric_columns": {
                    "vacancy_rate_pct": "float",
                    "special_drive_target_vacancies": "float",
                },
                "integer_columns": [
                    "sanctioned_strength",
                    "in_position_strength",
                    "vacancy_count",
                    "backlog_vacancies",
                ],
                "boolean_columns": [
                    "special_drive_flag",
                ],
                "string_columns": [
                    "sanction_id",
                    "fiscal_year",
                    "wc_id",
                    "discipline_id",
                    "grade_id",
                    "grade_band",
                    "social_category",
                    "gender",
                    "pwd_category",
                    "age_band",
                    "tenure_band",
                    "special_drive_name",
                ],
                "primary_key": ["sanction_id"],
            },
            "dpe_mou_scorecard.xlsx": {
                "target_table": "mou_targets",
                "sheet_strategy": {
                    "mode": "all_visible",
                    "exclude_sheets": ["_MoU_Grading_Rules"],
                },
                "header_row": 5,
                "unit_row": 6,
                "data_start_row": 7,
                "drop_columns": ["Unnamed: 0"],
                "subtotal_filter": {
                    "column": "mou_record_id",
                    "patterns": ["^Subtotal", "^Total", "^Summary"],
                },
                "footer_filter": {
                    "column": "mou_record_id",
                    "patterns": ["^Status:", "^Note:", "^Confidential"],
                },
                "numeric_columns": {
                    "excellent_target": "float",
                    "very_good_target": "float",
                    "good_target": "float",
                    "fair_target": "float",
                    "poor_target": "float",
                    "actual_achievement": "float",
                    "weighted_score": "float",
                    "composite_score": "float",
                    "projected_composite_score": "float",
                },
                "integer_columns": [
                    "weightage",
                ],
                "boolean_columns": [],
                "string_columns": [
                    "mou_record_id",
                    "fiscal_year",
                    "parameter_id",
                    "parameter_name",
                    "parameter_category",
                    "unit_of_measure",
                    "achieved_rating",
                    "notes_and_exceptions",
                ],
                "primary_key": ["mou_record_id"],
            },
            "csr_projects_tracker.xlsx": {
                "target_table": "csr_projects",
                "sheet_strategy": {
                    "mode": "all_visible",
                    "exclude_sheets": ["_Schedule_VII_Themes"],
                },
                "header_row": 5,
                "unit_row": 6,
                "data_start_row": 7,
                "drop_columns": ["Unnamed: 0"],
                "subtotal_filter": {
                    "column": "csr_project_id",
                    "patterns": ["^Subtotal", "^Total", "^Summary"],
                },
                "footer_filter": {
                    "column": "csr_project_id",
                    "patterns": ["^Mandate Note:", "^Note:", "^Confidential"],
                },
                "numeric_columns": {
                    "sanctioned_budget_inr_cr": "float",
                    "actual_disbursed_inr_cr": "float",
                },
                "integer_columns": [
                    "beneficiaries_count",
                ],
                "boolean_columns": [
                    "is_aspirational_district",
                    "is_operational_field_area",
                    "statutory_2pct_mandate_met_flag",
                ],
                "string_columns": [
                    "csr_project_id",
                    "project_name",
                    "fiscal_year",
                    "state_code",
                    "district_name",
                    "wc_id",
                    "csr_theme",
                    "implementing_agency",
                ],
                "primary_key": ["csr_project_id"],
            },
            "succession_planning.xlsx": {
                "target_table": "succession_plan",
                "sheet_strategy": {
                    "mode": "named",
                    "sheet_name": "succession_plan",
                    "exclude_sheets": ["_Talent_Bench_Scale"],
                },
                "header_row": 5,
                "unit_row": 6,
                "data_start_row": 7,
                "drop_columns": ["Unnamed: 0"],
                "subtotal_filter": {
                    "column": "position_id",
                    "patterns": ["^Subtotal", "^Total", "^Summary"],
                },
                "footer_filter": {
                    "column": "position_id",
                    "patterns": ["^Strictly Confidential", "^Note:"],
                },
                "numeric_columns": {},
                "integer_columns": [],
                "boolean_columns": [
                    "successor_1_has_ne_hard_station",
                    "successor_2_has_ne_hard_station",
                ],
                "string_columns": [
                    "position_id",
                    "role_title",
                    "wc_id",
                    "discipline_id",
                    "criticality_tier",
                    "incumbent_emp_id",
                    "incumbent_retire_date",
                    "fiscal_year",
                    "successor_1_emp_id",
                    "successor_1_readiness",
                    "successor_2_emp_id",
                    "successor_2_readiness",
                    "readiness_gap_reason",
                ],
                "primary_key": ["position_id"],
            },
            "legal_cases_tracker.xlsx": {
                "target_table": "grievances_litigation",
                "sheet_strategy": {
                    "mode": "named",
                    "sheet_name": "legal_cases",
                    "exclude_sheets": ["_Court_Jurisdiction_Codes"],
                },
                "header_row": 5,
                "unit_row": 6,
                "data_start_row": 7,
                "drop_columns": ["Unnamed: 0"],
                "subtotal_filter": {
                    "column": "case_id",
                    "patterns": ["^Subtotal", "^Total", "^Summary"],
                },
                "footer_filter": {
                    "column": "case_id",
                    "patterns": ["^Privileged", "^Note:"],
                },
                "numeric_columns": {
                    "age_years": "float",
                },
                "integer_columns": [
                    "age_days",
                ],
                "boolean_columns": [
                    "interim_stay_granted",
                    "is_aged_over_3_years",
                ],
                "string_columns": [
                    "case_id",
                    "emp_id",
                    "wc_id",
                    "case_type",
                    "case_category",
                    "forum",
                    "filing_date",
                    "fiscal_year",
                    "status",
                    "next_hearing_date",
                    "subject_summary",
                ],
                "primary_key": ["case_id"],
            },
        },
    }

    yaml_path = LANDING_EXCEL / "_parsing_specs.yaml"
    with open(yaml_path, "w", encoding="utf-8") as f:
        yaml.dump(specs, f, sort_keys=False, default_flow_style=False)
    print(f"Generated {yaml_path.name}")


def main() -> None:
    """Execute step 3.5 messy excel rewriting."""
    targets = [
        "manpower_plan.xlsx",
        "capex_be_re_actuals.xlsx",
        "dpe_mou_scorecard.xlsx",
        "csr_projects_tracker.xlsx",
        "succession_planning.xlsx",
        "legal_cases_tracker.xlsx",
    ]
    ensure_clean_backups(targets)

    print("Rewriting capex_be_re_actuals.xlsx...")
    process_capex_be_re_actuals()

    print("Rewriting manpower_plan.xlsx...")
    process_manpower_plan()

    print("Rewriting dpe_mou_scorecard.xlsx...")
    process_dpe_mou_scorecard()

    print("Rewriting csr_projects_tracker.xlsx...")
    process_csr_projects_tracker()

    print("Rewriting succession_planning.xlsx...")
    process_succession_planning()

    print("Rewriting legal_cases_tracker.xlsx...")
    process_legal_cases_tracker()

    print("Generating _parsing_specs.yaml...")
    generate_parsing_specs()
    print("Step 3.5 messy excel generation complete.")


if __name__ == "__main__":
    main()
