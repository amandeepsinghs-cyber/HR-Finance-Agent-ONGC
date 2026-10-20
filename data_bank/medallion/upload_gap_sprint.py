#!/usr/bin/env python3
"""medallion/upload_gap_sprint.py: Upload and load 4 gap-sprint files into Bronze (Step 4.1 & 4.2 extension).

Files:
1. landing/db_exports/env_incident_log.csv -> hrfin_bronze.db_exports_env_incident_log (482 rows)
2. landing/excel/posh_register.xlsx -> hrfin_bronze.excel_posh_register (58 rows)
3. landing/sap/apprentice_ledger.csv -> hrfin_bronze.sap_apprentice_ledger (7,585 rows)
4. landing/sap/EKKO_channel.csv -> hrfin_bronze.sap_ekko_channel (110,000 rows)

Batch date: 20261020
Preserves existing 42 tables and updates upload_manifest.csv & bronze_tables.yaml.
"""

from __future__ import annotations

import base64
import csv
import datetime
import hashlib
import io
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import yaml
from google.cloud import bigquery, storage

PROJECT_ID = "ongc-hr-finance-agent"
DATASET_ID = "hrfin_bronze"
LANDING_BUCKET = "ongc-hr-finance-agent-landing"
BATCH_DATE = "20261020"

BASE_DIR = Path(__file__).resolve().parent.parent
MANIFEST_PATH = BASE_DIR / "medallion" / "upload_manifest.csv"
BRONZE_YAML_PATH = BASE_DIR / "medallion" / "bronze" / "bronze_tables.yaml"

JOB_LABELS = {"datacloud": "ai-agent"}

GAP_FILES = [
    {
        "source": "db_exports",
        "file_name": "env_incident_log.csv",
        "local_path": "landing/db_exports/env_incident_log.csv",
        "table_name": "db_exports_env_incident_log",
        "is_excel": False,
    },
    {
        "source": "excel",
        "file_name": "posh_register.xlsx",
        "local_path": "landing/excel/posh_register.xlsx",
        "table_name": "excel_posh_register",
        "is_excel": True,
    },
    {
        "source": "sap",
        "file_name": "apprentice_ledger.csv",
        "local_path": "landing/sap/apprentice_ledger.csv",
        "table_name": "sap_apprentice_ledger",
        "is_excel": False,
    },
    {
        "source": "sap",
        "file_name": "EKKO_channel.csv",
        "local_path": "landing/sap/EKKO_channel.csv",
        "table_name": "sap_ekko_channel",
        "is_excel": False,
    },
]


def calculate_md5_and_size(file_path: Path):
    h = hashlib.md5()
    size = 0
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
            size += len(chunk)
    return h.hexdigest(), size


def upload_gap_files(gcs_client: storage.Client) -> List[Dict[str, str]]:
    bucket = gcs_client.bucket(LANDING_BUCKET)
    manifest_records = []

    for item in GAP_FILES:
        local_path = BASE_DIR / item["local_path"]
        blob_name = f"{item['source']}/{BATCH_DATE}/{item['file_name']}"
        gs_uri = f"gs://{LANDING_BUCKET}/{blob_name}"
        md5_hex, size = calculate_md5_and_size(local_path)

        blob = bucket.blob(blob_name)
        print(f"Uploading {local_path} -> {gs_uri}...")
        blob.upload_from_filename(str(local_path))
        blob.reload()

        remote_md5 = base64.b64decode(blob.md5_hash).hex()
        assert remote_md5 == md5_hex, f"MD5 mismatch for {blob_name}: {remote_md5} != {md5_hex}"
        assert blob.size == size, f"Size mismatch for {blob_name}: {blob.size} != {size}"
        print(f"  Verified {gs_uri} OK.")

        manifest_records.append({
            "local_path": item["local_path"],
            "gs_uri": gs_uri,
            "size": str(size),
            "md5": md5_hex,
        })

    return manifest_records


def update_manifest(new_records: List[Dict[str, str]]):
    with open(MANIFEST_PATH, "r", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    existing_uris = {r["gs_uri"] for r in rows}
    added = 0
    for rec in new_records:
        if rec["gs_uri"] not in existing_uris:
            rows.append(rec)
            added += 1

    rows_sorted = sorted(rows, key=lambda x: str(x["local_path"]))
    with open(MANIFEST_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["local_path", "gs_uri", "size", "md5"])
        writer.writeheader()
        writer.writerows(rows_sorted)

    print(f"Updated {MANIFEST_PATH}: added {added}, total {len(rows_sorted)} items.")


def load_csv_bronze(bq_client: bigquery.Client, item: Dict[str, Any], ingest_ts: datetime.datetime):
    local_path = BASE_DIR / item["local_path"]
    table_name = item["table_name"]
    staging_id = f"{PROJECT_ID}.{DATASET_ID}._staging_{table_name}"
    final_id = f"{PROJECT_ID}.{DATASET_ID}.{table_name}"
    gcs_uri = f"gs://{LANDING_BUCKET}/{item['source']}/{BATCH_DATE}/{item['file_name']}"

    with open(local_path, "r", encoding="utf-8", errors="replace") as fh:
        reader = csv.reader(fh)
        header = next(reader)
        expected_rows = sum(1 for _ in reader)

    try:
        schema = [bigquery.SchemaField(col, "STRING") for col in header]
        load_config = bigquery.LoadJobConfig(
            schema=schema,
            skip_leading_rows=1,
            source_format=bigquery.SourceFormat.CSV,
            write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
            allow_quoted_newlines=True,
            encoding="UTF-8",
            labels=JOB_LABELS,
        )
        job = bq_client.load_table_from_uri(gcs_uri, staging_id, job_config=load_config)
        job.result()

        ts_iso = ingest_ts.isoformat()
        query = f"""
        CREATE OR REPLACE TABLE `{final_id}` AS
        SELECT
          *,
          "{gcs_uri}" AS _source_file,
          TIMESTAMP("{ts_iso}") AS _ingest_ts,
          "{BATCH_DATE}" AS _batch_id
        FROM `{staging_id}`
        """
        q_job = bq_client.query(query, job_config=bigquery.QueryJobConfig(labels=JOB_LABELS))
        q_job.result()

        tbl = bq_client.get_table(final_id)
        assert tbl.num_rows == expected_rows, f"Row count mismatch for {table_name}: expected {expected_rows}, got {tbl.num_rows}"
        print(f"  [OK] {table_name}: {tbl.num_rows} rows, {len(tbl.schema)} cols")
        return {
            "source_type": item["source"],
            "source_file": gcs_uri,
            "local_path": item["local_path"],
            "row_count": tbl.num_rows,
            "column_count": len(tbl.schema),
            "columns": [f.name for f in tbl.schema],
        }
    finally:
        bq_client.delete_table(staging_id, not_found_ok=True)


def load_excel_bronze(bq_client: bigquery.Client, item: Dict[str, Any], ingest_ts: datetime.datetime):
    local_path = BASE_DIR / item["local_path"]
    table_name = item["table_name"]
    final_id = f"{PROJECT_ID}.{DATASET_ID}.{table_name}"
    gcs_uri = f"gs://{LANDING_BUCKET}/{item['source']}/{BATCH_DATE}/{item['file_name']}"

    wb = openpyxl.load_workbook(str(local_path), data_only=True)
    visible_sheets = [s for s in wb.sheetnames if wb[s].sheet_state != "hidden"]

    rows_data: List[Dict[str, Any]] = []
    max_cols = 0
    for sname in visible_sheets:
        ws = wb[sname]
        for row_idx, r in enumerate(ws.iter_rows(values_only=True), start=1):
            if not any(c is not None and str(c).strip() != "" for c in r):
                continue
            max_cols = max(max_cols, len(r))
            row_dict: Dict[str, Any] = {
                "_sheet_name": sname,
                "_row_number": row_idx,
            }
            for col_idx, c in enumerate(r):
                row_dict[f"col_{col_idx}"] = None if c is None else str(c)
            rows_data.append(row_dict)
    wb.close()

    arrow_fields = [
        pa.field("_sheet_name", pa.string()),
        pa.field("_row_number", pa.int64()),
    ]
    for c_idx in range(max_cols):
        arrow_fields.append(pa.field(f"col_{c_idx}", pa.string()))
    arrow_fields.extend([
        pa.field("_source_file", pa.string()),
        pa.field("_ingest_ts", pa.timestamp("us", tz="UTC")),
        pa.field("_batch_id", pa.string()),
    ])
    arrow_schema = pa.schema(arrow_fields)

    col_arrays = {f.name: [] for f in arrow_schema}
    for rd in rows_data:
        col_arrays["_sheet_name"].append(rd["_sheet_name"])
        col_arrays["_row_number"].append(rd["_row_number"])
        for c_idx in range(max_cols):
            col_arrays[f"col_{c_idx}"].append(rd.get(f"col_{c_idx}"))
        col_arrays["_source_file"].append(gcs_uri)
        col_arrays["_ingest_ts"].append(ingest_ts)
        col_arrays["_batch_id"].append(BATCH_DATE)

    arrow_table = pa.Table.from_pydict(col_arrays, schema=arrow_schema)
    buf = io.BytesIO()
    pq.write_table(arrow_table, buf)
    buf.seek(0)

    load_config = bigquery.LoadJobConfig(
        source_format=bigquery.SourceFormat.PARQUET,
        write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
        labels=JOB_LABELS,
    )
    job = bq_client.load_table_from_file(buf, final_id, job_config=load_config)
    job.result()

    tbl = bq_client.get_table(final_id)
    assert tbl.num_rows == len(rows_data), f"Row count mismatch for {table_name}: expected {len(rows_data)}, got {tbl.num_rows}"
    print(f"  [OK] {table_name}: {tbl.num_rows} rows, {len(tbl.schema)} cols")
    return {
        "source_type": item["source"],
        "source_file": gcs_uri,
        "local_path": item["local_path"],
        "row_count": tbl.num_rows,
        "column_count": len(tbl.schema),
        "columns": [f.name for f in tbl.schema],
    }


def update_bronze_yaml(new_table_meta: Dict[str, Any]):
    with open(BRONZE_YAML_PATH, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)

    for tname, meta in new_table_meta.items():
        data["tables"][tname] = meta

    data["total_tables"] = len(data["tables"])
    data["total_rows"] = sum(m["row_count"] for m in data["tables"].values())

    with open(BRONZE_YAML_PATH, "w", encoding="utf-8") as f:
        yaml.dump(data, f, sort_keys=False, default_flow_style=False)

    print(f"Updated {BRONZE_YAML_PATH}: total tables={data['total_tables']}, total rows={data['total_rows']}")


def main():
    print("Connecting to GCS and BigQuery...")
    gcs_client = storage.Client(project=PROJECT_ID)
    bq_client = bigquery.Client(project=PROJECT_ID)

    print("Uploading 4 gap-sprint files to GCS...")
    new_manifest_records = upload_gap_files(gcs_client)
    update_manifest(new_manifest_records)

    print("Loading 4 gap-sprint tables into BigQuery hrfin_bronze...")
    ingest_ts = datetime.datetime.now(datetime.timezone.utc)
    new_tables_meta = {}

    for item in GAP_FILES:
        if item["is_excel"]:
            meta = load_excel_bronze(bq_client, item, ingest_ts)
        else:
            meta = load_csv_bronze(bq_client, item, ingest_ts)
        new_tables_meta[item["table_name"]] = meta

    update_bronze_yaml(new_tables_meta)
    print("Gap sprint files successfully uploaded and ingested into Bronze!")


if __name__ == "__main__":
    main()
